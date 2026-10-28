"""SENER field-production workbook reader, independent of browser automation.

Source: SENER, Producción nacional de petróleo y gas natural, enero 2025,
https://hidrocarburos.energia.gob.mx/estadisticas/; units and headers in each
sheet. No BOE conversion, liquids substitution or missing-to-zero imputation.
"""

from __future__ import annotations

import calendar
import hashlib
import math
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

LISTING_URL = "https://hidrocarburos.energia.gob.mx/estadisticas/"
KEYS = ["date", "basin", "location", "field_name", "operator"]
LABELS = ["Fecha", "Cuenca", "Ubicación", "Campo", "Operador"]
SHEETS = {
    "Aceite y condensado": {
        "Petróleo (Mbd)": "oil_bbl",
        "Condensado (Mbd)": "condensate_bbl",
    },
    "Gas Natural": {"Gas natural sin nitrógeno (MMpcd)": "gas_mcf"},
}
# SENER January 2025 workbook, units row: Mbd = thousand bbl/day;
# MMpcd = million cubic feet/day = thousand Mcf/day. Both multipliers = 1000.
RATE_TO_MONTHLY = 1000


@dataclass(frozen=True)
class ProductionWorkbook:
    """Observed monthly volumes and integrity evidence, not an economic forecast."""

    production: pd.DataFrame
    sha256: str
    vintage: str


def _volume(rate, date):
    if rate is None:
        return float("nan")
    if isinstance(rate, bool) or not isinstance(rate, (int, float)):
        raise TypeError("production rates must be numeric or blank")
    if not math.isfinite(rate) or rate < 0:
        raise ValueError("production rates must be finite and non-negative")
    # SENER publishes monthly average daily rates; calendar.monthrange supplies
    # Gregorian month length, including leap years, with no assumed 30-day month.
    volume = rate * RATE_TO_MONTHLY * calendar.monthrange(date.year, date.month)[1]
    if not math.isfinite(volume):
        raise ValueError("monthly volume must be finite")
    return volume


def _read_sheet(sheet, metrics):
    rows = iter(sheet.iter_rows(values_only=True))
    for header in rows:
        if list(header[: len(LABELS)]) == LABELS:
            break
    else:
        raise ValueError(f"missing field-production header in {sheet.title}")
    if not set(metrics).issubset(header):
        raise ValueError(
            f"missing production columns or changed units in {sheet.title}"
        )
    indices = {label: header.index(label) for label in metrics}
    records = []
    for row in rows:
        if all(value is None for value in row):
            continue
        date, *identity = row[: len(KEYS)]
        if not isinstance(date, datetime) or date.day != 1:
            raise ValueError("monthly production date must be the first day of a month")
        if any(not isinstance(v, str) or not v.strip() for v in identity):
            raise ValueError("field-production identity must be non-empty text")
        record = dict(zip(KEYS, [date, *[v.strip() for v in identity]]))
        record.update(
            {
                target: _volume(row[indices[label]], date)
                for label, target in metrics.items()
            }
        )
        records.append(record)
    frame = pd.DataFrame(records, columns=[*KEYS, *metrics.values()])
    if frame.duplicated(KEYS).any():
        raise ValueError(f"duplicate field-month identity in {sheet.title}")
    return frame


def load_production_workbook(path: str | Path) -> ProductionWorkbook:
    """Read both published sheets; reject schema drift and preserve unknowns.

    Oil excludes condensate. Gas excludes nitrogen. Water is not published.
    Fields with different basin, location or operator remain separate records.
    """
    path = Path(path)
    payload = path.read_bytes()
    digest = hashlib.sha256(payload).hexdigest()
    book = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        frames = [_read_sheet(book[name], metrics) for name, metrics in SHEETS.items()]
    except KeyError as exc:
        raise ValueError("required field-production sheet missing") from exc
    finally:
        book.close()
    frame = frames[0].merge(frames[1], on=KEYS, how="outer", validate="one_to_one")
    if frame.empty:
        raise ValueError("workbook contains no field-production records")
    frame["year"] = frame.date.dt.year
    frame["month"] = frame.date.dt.month
    frame["water_bbl"] = float("nan")
    frame["region"] = "mexico"
    frame["source"] = "sener_field_production"
    return ProductionWorkbook(frame, digest, frame.date.max().strftime("%Y-%m"))
