"""Offline adapters for reviewed OPRED pages and public HSE spreadsheets.

No network request is made here. Obtain public inputs only after checking access
and terms. The existing offshore model, citation builder and transactional intake
remain authoritative; the Gulf-only HSEIncident is never used.

Layouts: OPRED PON1 notifications 4966 and 6788 (2024), read 2026-10-22.
HSE spreadsheet layout is NOT_EVALUATED: its robots rule blocks acquisition.
Column mapping, worksheet, header row and date format therefore require inputs.
Rules: issue #1325 (2026), Fields, Deduplication, Citation and Never collect.
"""

from __future__ import annotations

import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from html.parser import HTMLParser
from pathlib import Path
from typing import Literal
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from pydantic import ValidationError

from worldenergydata.hse.offshore import Event, SourceDocument, check_privacy

OGL_URL = "https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/"
HSE_SPREADSHEET_URL = (
    "https://www.hse.gov.uk/enforce/assets/docs/temporary-register-notices.xlsx"
)
HSE_MANUAL_URLS = (
    "https://www.hse.gov.uk/enforce/temporary-register-notices.htm",
    HSE_SPREADSHEET_URL,
    "https://notices.hse.gov.uk/",
    "https://resources.hse.gov.uk/notices",
)
# HSE temporary register landing page, Details and OGL footer (2026-10-22).
# OGL v3.0, Exemptions: personal data are excluded from the licence.
PERSON_COLUMN = re.compile(
    r"inspector|witness|injured|deceased|person|contact|e-?mail|telephone|phone",
    re.IGNORECASE,
)


class _PON1Page(HTMLParser):
    """Read row headers/cells in main, without retaining prose or navigation."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.in_main = False
        self.in_footer = False
        self.cell = None
        self.row = []
        self.rows = []
        self.heading = None
        self.headings = []
        self.licence_links = []
        self.footer = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "main":
            self.in_main = True
        if tag == "footer":
            self.in_footer = True
        if self.in_footer and tag == "a" and attrs.get("rel") == "license":
            self.licence_links.append(attrs.get("href"))
        if self.in_main:
            if tag == "tr":
                self.row = []
            elif tag in {"th", "td"}:
                self.cell = (tag, [])
            elif tag == "h1":
                self.heading = []
            elif tag == "br" and self.cell is not None:
                self.cell[1].append(" ")

    def handle_data(self, data):
        if self.cell is not None:
            self.cell[1].append(data)
        if self.heading is not None:
            self.heading.append(data)
        if self.in_footer:
            self.footer.append(data)

    def handle_endtag(self, tag):
        if tag in {"th", "td"} and self.cell is not None:
            self.row.append((self.cell[0], " ".join("".join(self.cell[1]).split())))
            self.cell = None
        elif tag == "tr" and self.in_main:
            if [kind for kind, _ in self.row] != ["th", "td"]:
                raise ValueError("PON1 layout: expected one header and one value")
            self.rows.append((self.row[0][1], self.row[1][1]))
        elif tag == "h1" and self.heading is not None:
            self.headings.append(" ".join("".join(self.heading).split()))
            self.heading = None
        elif tag == "main":
            self.in_main = False
        elif tag == "footer":
            self.in_footer = False


def parse_pon1(
    html: str,
    *,
    context: Mapping,
    url: str,
    access_date: date,
    reviewed_text: Sequence[str],
) -> dict:
    """Build one reviewed evidence packet; never guess geography or a cause.

    Context supplies only curated classification, country/area and aliases.
    Published date, installation and release facts cannot be overridden. The
    raw and normalized text inventories both need explicit privacy review.
    Multiple substances are refused because Event has a single release tuple.
    No unit conversion or coordinate conversion is performed.
    """
    context_fields = {
        "event_id",
        "country",
        "basin_or_area",
        "operation",
        "equipment",
        "failure_mode",
        "failure_summary",
        "installation_aliases",
    }
    if set(context) - context_fields:
        raise ValueError("PON1 context contains published or unsupported fields")
    page = _PON1Page()
    page.feed(html)
    page.close()
    if (
        OGL_URL not in page.licence_links
        or "Open Government Licence v3.0" not in " ".join(page.footer)
    ):
        raise ValueError("PON1 licence not established")
    if len(page.headings) != 1 or not re.fullmatch(
        r"IRS/\d{4}/\d+/PON1", page.headings[0]
    ):
        raise ValueError("PON1 layout: notification reference missing")
    fields = dict(page.rows)
    if len(fields) != len(page.rows):
        raise ValueError("PON1 layout: repeated field or multiple releases")
    check_privacy([list(fields.values()), page.headings, dict(context)], reviewed_text)
    required = {
        "Status",
        "First observed date",
        "Incident type(s)",
        "Installation",
        "Source",
    }
    if not required <= fields.keys() or any(not fields[key] for key in required):
        raise ValueError("PON1 layout: required field missing")
    release_type = fields["Incident type(s)"]
    if release_type == "Chemical release":
        product_key = "Chemical product released"
        quantity_key = "Total quantity of chemical released"
    elif release_type == "Oil release":
        product_key = "Oil product released"
        quantity_key = "Total quantity of oil released"
    else:
        raise ValueError("PON1 layout: mixed or unsupported release types")
    other_product = (
        "Oil product released"
        if release_type == "Chemical release"
        else "Chemical product released"
    )
    if other_product in fields or not fields.get(product_key):
        raise ValueError("PON1 layout: ambiguous or missing release substance")
    quantity = re.fullmatch(r"(\d+(?:\.\d+)?)\s+(\S+)", fields.get(quantity_key, ""))
    if quantity is None:
        raise ValueError("PON1 quantity must include one nonnegative number and unit")
    try:
        # OPRED First observed date, notifications 4966/6788 (2024).
        # No UTC offset is published for First observed time: keep date precision.
        observed = datetime.strptime(fields["First observed date"], "%d %B %Y").date()
        event = Event(
            **context,
            event_date=observed,
            installation_name=fields["Installation"],
            installation_type=fields.get("Installation type") or None,
            field=fields.get("Field") or None,
            block_or_lease=re.sub(
                r"\s*/\s*", "/", fields.get("Quadrant / block no.", "")
            )
            or None,
            release_substance=fields[product_key],
            release_quantity=Decimal(quantity[1]),
            release_unit=quantity[2],
            causes_as_stated=[fields["Source"]],
            organisations=[
                part.strip()
                for part in fields.get("Organisations involved", "").split(",")
                if part.strip()
            ],
        )
        source = SourceDocument(
            event_id=event.event_id,
            source="opred",
            source_ref=page.headings[0],
            title=page.headings[0],
            doc_type="release notice",
            issue_date=None,
            status=fields["Status"],
            url=url,
            access_date=access_date,
            licence="OGL-3.0",
            evidence_level="regulator",
            access_status="public",
        )
    except (ValueError, TypeError, ValidationError):
        # Validation diagnostics must not echo raw values or locations.
        raise ValueError("PON1 fields failed validation") from None
    packet = {
        "events": [event.model_dump(mode="json")],
        "sources": [source.model_dump(mode="json")],
        "reviewed_text": list(reviewed_text),
        "provenance_notes": [],
    }
    if fields.get("First observed time"):
        packet["provenance_notes"].append(
            f"Observed time {fields['First observed time']} has no published timezone; date only."
        )
    if fields.get("Formal investigation ever opened"):
        packet["provenance_notes"].append(
            f"Formal investigation ever opened: {fields['Formal investigation ever opened']}."
        )
    check_privacy(
        {"events": packet["events"], "sources": packet["sources"]}, reviewed_text
    )
    return packet


@dataclass(frozen=True)
class NoticeLoadResult:
    status: Literal["ready", "manual_required"]
    sources: tuple[SourceDocument, ...] = ()
    manual_urls: tuple[str, ...] = ()
    reason: str | None = None


def manual_hse_step(
    reason: Literal["BLOCKED-NETWORK", "FILE_UNREADABLE", "READER_UNAVAILABLE"],
) -> NoticeLoadResult:
    """Return official URLs for a human step, with no automatic retry or bypass."""
    return NoticeLoadResult(
        "manual_required", manual_urls=HSE_MANUAL_URLS, reason=reason
    )


def _notice_date(value, date_format):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(value, date_format).date()


def _notice_ref(value):
    if (
        isinstance(value, bool)
        or not isinstance(value, (str, int))
        or not str(value).strip()
    ):
        raise ValueError("HSE reference must be text or an integer")
    return str(value).strip()


def load_hse_notices(
    path: Path,
    *,
    columns: Mapping[str, str],
    sheet: str,
    header_row: int,
    associations: Mapping[str, str],
    date_format: str,
    access_date: date,
    reviewed_text: Sequence[str],
) -> NoticeLoadResult:
    """Map manually obtained public XLSX rows to explicitly associated events.

    A notice issue date is never used as the incident date. Only references in
    associations are selected, and all selected row strings require review,
    including unused columns. Populated personal columns are always refused.
    No formula is executed or replaced by a potentially stale cached value.
    HSE's blocked live layout must be checked manually before choosing columns.
    """
    required = {"source_ref", "title", "doc_type", "issue_date", "status"}
    optional = {"compliance_date", "legal_provision", "url", "summary"}
    # openpyxl 3.1 Worksheet.iter_rows, min_row: one-based index
    # (official worksheet API, accessed 2026-10-22).
    if (
        not required <= columns.keys()
        or set(columns) - (required | optional)
        or len(set(columns.values())) != len(columns)
        or type(header_row) is not int
        or header_row < 1
    ):
        raise ValueError("HSE layout mapping is invalid")
    if not associations:
        raise ValueError("HSE association references required")
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        return manual_hse_step("FILE_UNREADABLE")
    try:
        from openpyxl import load_workbook
    except ImportError:
        return manual_hse_step("READER_UNAVAILABLE")
    try:
        book = load_workbook(path, read_only=True, data_only=False)
    except (OSError, BadZipFile, ValueError, KeyError, ParseError, EOFError):
        return manual_hse_step("FILE_UNREADABLE")
    try:
        if sheet not in book.sheetnames:
            raise ValueError("HSE layout: worksheet missing")
        rows = book[sheet].iter_rows(min_row=header_row)
        header_cells = next(rows, ())
        headers = [cell.value for cell in header_cells]
        nonempty = [value for value in headers if value is not None]
        if len(set(nonempty)) != len(nonempty) or not set(columns.values()) <= set(
            headers
        ):
            raise ValueError("HSE layout: missing or duplicate column headers")
        indexes = {key: headers.index(label) for key, label in columns.items()}
        selected = []
        seen = set()
        for cells in rows:
            values = [cell.value for cell in cells]
            value = values[indexes["source_ref"]]
            if value is None:
                continue
            reference = _notice_ref(value)
            if reference not in associations:
                continue
            if reference in seen:
                raise ValueError("HSE duplicate reference in spreadsheet")
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError("HSE formula requires manual review")
            if any(
                isinstance(label, str)
                and PERSON_COLUMN.search(label)
                and values[index] not in (None, "")
                for index, label in enumerate(headers)
            ):
                raise ValueError("privacy: personal column; admission refused")
            check_privacy(values, reviewed_text)
            row = {key: values[index] for key, index in indexes.items()}
            try:
                source = SourceDocument(
                    event_id=associations[reference],
                    source="uk_hse",
                    source_ref=reference,
                    title=row["title"],
                    doc_type=row["doc_type"],
                    issue_date=_notice_date(row["issue_date"], date_format),
                    compliance_date=_notice_date(
                        row.get("compliance_date"), date_format
                    ),
                    legal_provision=row.get("legal_provision") or None,
                    status=row["status"],
                    url=row.get("url") or HSE_SPREADSHEET_URL,
                    access_date=access_date,
                    licence="OGL-3.0",
                    evidence_level="regulator",
                    access_status="public",
                    summary=row.get("summary") or None,
                    association_basis="manual association; notice issue date is not the event date",
                )
                if source.doc_type not in {"improvement notice", "prohibition notice"}:
                    raise ValueError("unsupported notice type")
            except (ValueError, TypeError, ValidationError):
                raise ValueError("HSE fields failed validation") from None
            check_privacy(source.model_dump(mode="json"), reviewed_text)
            selected.append(source)
            seen.add(reference)
        if seen != set(associations):
            raise ValueError("HSE association references not found")
        return NoticeLoadResult("ready", sources=tuple(selected))
    except (OSError, BadZipFile, ParseError, EOFError):
        return manual_hse_step("FILE_UNREADABLE")
    finally:
        book.close()
