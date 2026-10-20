"""Read-only workbook bridges and deterministic integrated decision surface.

Cell coordinates and category definitions: frozen financial_project_summary.xlsx,
Project_Summary!E1:M1,P1:Q1 and project sheet headers F1:G1,I1,K1:L1;
lease_assumptions.xlsx, assumptions!B1:F1 and row 15 (FDAS V30).
These are workbook assumptions, not independently sourced disclosures.
"""

import csv
import hashlib
import io
import json
from collections import defaultdict
from decimal import Decimal, InvalidOperation, localcontext
from html import escape
from pathlib import Path

from .models import Basis, Interval, MaterialityPolicy, Observation
from .reconcile import reconcile
from .workbook_reader import canonical, digest, read_frozen, validate_crosswalk

FDAS = Path("docs/modules/bsee/analysis/production/FDAS_V30")
CONTRACT = Path("config/analysis/lower_tertiary/fdas_workbook_vintages.json")
CURATED = Path("data/modules/cost/curated")
# The workbook cost units are USD; assumption row 15 is in millions USD.
# NIST, Metric (SI) Prefixes, SI Prefixes table, mega (accessed 2026-10-19):
# https://www.nist.gov/pml/owm/metric-si-prefixes
MILLION = Decimal(1000000)
SUMMARY = {
    "E": "host",
    "F": "surf",
    "G": "booster_pump",
    "H": "injection_pump",
    "I": "dry_system",
    "J": "facilities",
    "K": "drilling",
    "L": "completion",
    "M": "dc_total",
    "P": "variable_opex",
    "Q": "fixed_opex",
}
MONTHLY = {
    "F": "variable_opex",
    "G": "fixed_opex",
    "I": "drilling",
    "K": "completion",
    "L": "facilities",
}
# Unit denominators are stated in assumptions!G8:G17,G36:G38, V30.
PARAMETER_UNITS = {
    8: "USD MM/day",
    9: "USD MM/day",
    10: "USD MM",
    12: "USD MM/well",
    13: "USD MM/producer",
    15: "USD MM/facility",
    16: "USD/bbl",
    17: "USD MM/year",
    36: "USD MM/pump",
    37: "USD MM/pump",
    38: "USD MM/pump",
}
UPSTREAM = [
    "data/modules/cost/derived/cost_synthesis_manifest.v3.json",
    "data/modules/cost/derived/project_trace_contract_manifest.v3.json",
    "data/modules/cost/derived/capex_estimator_manifest.json",
]
FIELDS = (
    "project_id",
    "component",
    "start",
    "end",
    "currency",
    "unit",
    "assumption_vintage",
    "provenance",
    "source",
    "disposition",
    "workbook_value",
    "monthly_value",
    "disclosed_value",
    "mapped_value",
    "modeled_value",
    "variance",
    "classification",
    "materiality",
    "status",
    "reason",
)


def _number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise TypeError(
            "required workbook numeric cell is missing or not a cached number"
        )
    result = Decimal(str(value))
    if not result.is_finite():
        raise ValueError("nonfinite workbook number")
    return result


def _sum(values):
    with localcontext() as context:
        context.prec = (
            80  # Exact workbook-scale decimal accumulation, no acceptance tolerance.
        )
        return sum(values, Decimal(0))


def _subtract(left, right):
    with localcontext() as context:
        context.prec = 80
        return left - right


def workbook_rows(root: Path):
    from openpyxl import load_workbook

    contract = json.loads((root / CONTRACT).read_text())
    inventory = {
        name: read_frozen(root / FDAS / name, profile)
        for name, profile in contract["workbooks"].items()
    }
    if any(
        sheet["uncached_formula_cells"]
        for book in inventory.values()
        for sheet in book["sheets"]
    ):
        raise ValueError("uncached formulas cannot participate in reconciliation")
    financial = load_workbook(
        root / FDAS / "financial_project_summary.xlsx", read_only=True, data_only=True
    )
    assumptions = load_workbook(
        root / FDAS / "lease_assumptions.xlsx", read_only=True, data_only=True
    )
    try:
        summary = financial.worksheets[0]
        systems = {
            cell.value: cell.column
            for cell in assumptions.worksheets[0][1]
            if cell.value in {"subsea15", "subsea20", "dry", "tieback15", "tieback20"}
        }
        parameters = []
        for column in systems.values():
            for parameter_row, cells in enumerate(
                assumptions.worksheets[0].iter_rows(
                    min_row=2, min_col=column, max_col=column
                ),
                2,
            ):
                cell = cells[0]
                parameters.append(
                    {
                        "source": f"lease_assumptions.xlsx:sheet-000!{chr(ord('A') + column - 1)}{parameter_row}",
                        "system": assumptions.worksheets[0].cell(1, column).value,
                        "parameter_id": f"parameter-{parameter_row:03d}",
                        "value": (
                            str(cell.value)
                            if isinstance(cell.value, (int, float, bool))
                            else None
                        ),
                        "text_value_sha256": (
                            hashlib.sha256(str(cell.value).encode()).hexdigest()
                            if cell.value is not None
                            and not isinstance(cell.value, (int, float, bool))
                            else None
                        ),
                        "unit": PARAMETER_UNITS.get(
                            parameter_row, "NOT_EVALUATED: non-cost parameter unit"
                        ),
                        "assumption_vintage": contract["assumption_vintage"],
                    }
                )
        rows, traces, identities, cell_crosswalk = [], [], [], []
        for index, summary_cells in enumerate(summary.iter_rows(min_row=2), 1):
            row_number = index + 1
            identity = summary_cells[0].value
            # Exact workbook-internal identity only; never fuzzy-match external projects.
            if identity not in financial.sheetnames[1:]:
                raise ValueError(
                    "summary project lacks an exact monthly sheet identity"
                )
            sheet = financial[identity]
            sheet_index = financial.sheetnames.index(identity)
            project_id = f"workbook-project-{index:03d}"
            month_rows = list(sheet.iter_rows(min_row=2, values_only=True))
            periods = [values[0].strftime("%Y-%m") for values in month_rows]
            if periods != sorted(set(periods)):
                raise ValueError("monthly periods must be unique and ordered")
            identities.append(
                {
                    "project_id": project_id,
                    "summary_row": row_number,
                    "sheet_id": f"sheet-{sheet_index:03d}",
                    "system": summary_cells[1].value,
                    "external_identity": "NOT_EVALUATED",
                }
            )
            monthly = {
                component: _sum(
                    _number(values[ord(column) - ord("A")]) for values in month_rows
                )
                for column, component in MONTHLY.items()
            }
            for period, values in zip(periods, month_rows):
                traces.append(
                    {
                        "project_id": project_id,
                        "period": period,
                        **{
                            component: str(_number(values[ord(column) - ord("A")]))
                            for column, component in MONTHLY.items()
                        },
                    }
                )
            values = {
                component: _number(summary[f"{column}{row_number}"].value)
                for column, component in SUMMARY.items()
            }
            visible_facilities = _sum(
                values[k]
                for k in [
                    "host",
                    "surf",
                    "booster_pump",
                    "injection_pump",
                    "dry_system",
                ]
            )
            values["injection_facility"] = _subtract(
                values["facilities"], visible_facilities
            )
            values["development_total"] = _sum(
                values[k] for k in ["facilities", "drilling", "completion"]
            )
            monthly["development_total"] = _sum(
                monthly[k] for k in ["facilities", "drilling", "completion"]
            )
            monthly["dc_total"] = _sum(monthly[k] for k in ["drilling", "completion"])
            crosswalk = [
                {
                    "cell": f"{column}{row_number}",
                    "component": component,
                    "disposition": (
                        "aggregate"
                        if component in {"facilities", "dc_total"}
                        else "included"
                    ),
                }
                for column, component in SUMMARY.items()
            ]
            validate_crosswalk({f"{c}{row_number}" for c in SUMMARY}, crosswalk)
            cell_crosswalk.extend(
                {"project_id": project_id, "sheet_id": "sheet-000", **entry}
                for entry in crosswalk
            )
            for component in [*values, "installation", "other"]:
                total = values.get(component)
                monthly_value = monthly.get(component)
                column = next(
                    (c for c, name in SUMMARY.items() if name == component), None
                )
                source = (
                    f"financial_project_summary.xlsx:sheet-000!{column}{row_number}"
                    if column
                    else f"financial_project_summary.xlsx:summary-row-{row_number}; derived or unavailable"
                )
                difference = (
                    _subtract(total, monthly_value)
                    if total is not None and monthly_value is not None
                    else None
                )
                reason = "no independent compatible comparison or approved materiality policy"
                if component == "injection_facility":
                    system_column = systems.get(summary_cells[1].value)
                    if system_column is None:
                        raise ValueError("unknown workbook development system")
                    facility_input = (
                        _number(assumptions.worksheets[0].cell(15, system_column).value)
                        * MILLION
                    )
                    reason = (
                        "embedded residual equals system facility input"
                        if total == facility_input
                        else "residual not established as system facility input"
                    )
                rows.append(
                    {
                        **dict.fromkeys(FIELDS, ""),
                        "project_id": project_id,
                        "component": component,
                        "start": periods[0],
                        "end": periods[-1],
                        "currency": "USD",
                        "unit": "USD",
                        "assumption_vintage": contract["assumption_vintage"],
                        "provenance": "workbook_assumption",
                        "source": source,
                        "disposition": (
                            "unmapped"
                            if total is None
                            else (
                                "aggregate"
                                if component
                                in {"facilities", "dc_total", "development_total"}
                                else "included"
                            )
                        ),
                        "workbook_value": str(total) if total is not None else "",
                        "monthly_value": (
                            str(monthly_value) if monthly_value is not None else ""
                        ),
                        "variance": str(difference) if difference is not None else "",
                        "classification": "missing_evidence",
                        "materiality": "NOT_EVALUATED",
                        "status": "NOT_EVALUATED",
                        "reason": reason,
                    }
                )
        inventory["parameters"] = parameters
        inventory["projects"] = identities
        inventory["cell_crosswalk"] = cell_crosswalk
        inventory["time_basis"] = contract["time_basis"]
        inventory["v30_reproduction"] = {
            "status": "NOT_EVALUATED",
            "reason": "monthly aggregation is an arithmetic check; financial engine reproduction is not established",
        }
        return rows, inventory, traces
    finally:
        financial.close()
        assumptions.close()


def _csv(rows, fields):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=fields, lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue()


def _table(rows, fields):
    headings = "".join(f"<th>{escape(field)}</th>" for field in fields)
    contents = "".join(
        "<tr>"
        + "".join(
            f"<td>{escape(str(row.get(field, '')) if field in {'period', 'start', 'end'} else _display(row.get(field, '')))}</td>"
            for field in fields
        )
        + "</tr>"
        for row in rows
    )
    return (
        f"<table><thead><tr>{headings}</tr></thead><tbody>{contents}</tbody></table>"
        if rows
        else "<p>NOT_EVALUATED — compatible evidence unavailable.</p>"
    )


def _display(value):
    # Display rounding only, per reporting convention; CSV/JSON retain native decimals.
    try:
        number = Decimal(str(value))
        if number.is_finite():
            return f"{number:,.1f}"
    except InvalidOperation:
        return str(value)
    return str(value)


def render_html(payload: dict) -> str:
    rows, assets = payload["rows"], payload["assets"]
    totals = [r for r in rows if r.get("component") == "development_total"]
    body = [
        "<!doctype html><html lang='en'><meta charset='utf-8'><title>FDAS cost reconciliation</title>",
        "<style>body{font:16px system-ui;margin:2rem;color:#17324a}table{border-collapse:collapse;display:block;overflow:auto;margin:1rem 0}td,th{padding:.4rem;border:1px solid #ccd}th{background:#eef}section{margin:2rem 0}</style>",
        "<h1>FDAS workbook reconciliation and cost map</h1>",
        "<p>Read-only workbook assumptions. Stakeholder circulation requires explicit maintainer approval under #1017. V50 differences remain under #899; no assumptions are adopted here.</p>",
        "<p>Summary: workbook-internal monthly and component bridges are shown separately from external cost evidence. Missing comparisons and materiality policy remain NOT_EVALUATED. Full V30 engine reproduction is NOT_EVALUATED.</p>",
        "<h2>Revision history</h2><table><tr><th>Revision</th><th>Date</th><th>Change and reason</th></tr><tr><td>1</td><td>2026-10-19</td><td>Initial read-only comparison to expose basis and evidence gaps.</td></tr></table>",
    ]
    sections = [
        (
            "Portfolio summary",
            _table(totals, ["project_id", "workbook_value", "monthly_value", "status"]),
        ),
        (
            "Project drill-down",
            _table(
                payload["inventory"].get("projects", []),
                [
                    "project_id",
                    "summary_row",
                    "sheet_id",
                    "system",
                    "external_identity",
                ],
            ),
        ),
        (
            "Asset to total",
            _table(
                [a for a in assets if a.get("direction") == "asset_to_project"],
                [
                    "project_id",
                    "requirement_id",
                    "row_kind",
                    "value_low_mm",
                    "value_high_mm",
                    "evidence_derivation",
                    "counting_disposition",
                ],
            ),
        ),
        (
            "Required assets and workbook links",
            _table(
                payload["inventory"].get("requirement_links", []),
                [
                    "project_id",
                    "requirement_ids",
                    "workbook",
                    "cell",
                    "mapping_status",
                    "unit",
                    "source",
                ],
            ),
        ),
        (
            "Total to asset",
            _table(
                [a for a in assets if a.get("direction") == "project_to_asset"],
                [
                    "project_id",
                    "requirement_id",
                    "award_id",
                    "row_kind",
                    "value_low_mm",
                    "value_high_mm",
                    "evidence_derivation",
                    "counting_disposition",
                ],
            ),
        ),
        (
            "Cost time traces",
            "<p>Workbook annual sums below retain their monthly source cells in monthly.csv. They are assumption cash flows, not historical disclosures.</p>"
            + _table(payload["traces"], ["project_id", "period", *MONTHLY.values()]),
        ),
        (
            "Estimator backtests",
            "<p>NOT_EVALUATED — estimator manifest and evidenced project holdouts are required; no error or prediction is invented.</p>",
        ),
        (
            "Workbook reconciliation",
            "<p>Variance here is summary minus monthly sum in native USD. No acceptance tolerance is inferred. Aggregate rows overlap their components and must not be summed again.</p>"
            + _table(rows, FIELDS),
        ),
        (
            "External comparisons",
            _table(
                payload["comparisons"],
                [
                    "project_id",
                    "component",
                    "status",
                    "reason",
                    "variance",
                    "materiality",
                    "classification",
                    "evidence",
                ],
            ),
        ),
        (
            "Provenance legend",
            "<p>workbook_assumption: frozen inputs/results, never disclosures. disclosed: independent public evidence. mapped: linked evidence. allocated/modeled: inference, never award disclosure. Empty is unavailable, not zero. NOT_EVALUATED means evidence, compatibility or approval is missing.</p>",
        ),
        (
            "Upstream contracts",
            _table(payload["upstream"], ["path", "sha256", "status", "reason"]),
        ),
    ]
    body.extend(
        f"<section><h2>{title}</h2>{content}</section>" for title, content in sections
    )
    return "\n".join(body) + "\n</html>\n"


def compare_rows(rows: list[dict], specification: dict) -> list[dict]:
    """Compare supplied evidence with frozen values, never a supplied reference value.

    The caller must document workbook basis equivalence. Approval and materiality
    are explicit inputs. The input digest is recorded by the builder.
    """
    from dataclasses import asdict

    policy_input = specification.get("policy")
    policy = (
        MaterialityPolicy(
            Decimal(policy_input["absolute"]),
            Decimal(policy_input["relative"]),
            policy_input["source"],
            policy_input["approved"],
        )
        if policy_input is not None
        else None
    )
    lookup = {(row["project_id"], row["component"]): row for row in rows}
    results, seen = [], set()
    for entry in specification["comparisons"]:
        key = (entry["project_id"], entry["component"])
        if key in seen:
            raise ValueError("duplicate external comparison")
        seen.add(key)
        row = lookup[key]
        basis = Basis(**entry["workbook_basis"])
        if any(
            getattr(basis, field) != row[field]
            for field in [
                "project_id",
                "component",
                "start",
                "end",
                "currency",
                "unit",
                "assumption_vintage",
            ]
        ):
            raise ValueError("supplied workbook basis disagrees with frozen row")
        if not row["workbook_value"] or not entry.get("basis_evidence", "").strip():
            results.append(
                {
                    "project_id": key[0],
                    "component": key[1],
                    "status": "NOT_EVALUATED",
                    "reason": "missing workbook value or basis evidence",
                }
            )
            continue
        native = Decimal(row["workbook_value"])
        reference = Observation(
            basis, Interval(native, native), "workbook_assumption", row["source"]
        )
        candidate = entry["comparison"]
        candidate_amount = candidate["amount"]
        other = Observation(
            Basis(**candidate["basis"]),
            Interval(
                *(
                    (
                        Decimal(candidate_amount[bound])
                        if candidate_amount[bound] is not None
                        else None
                    )
                    for bound in ["low", "high"]
                )
            ),
            candidate["provenance"],
            candidate["source"],
        )
        result = reconcile(
            reference,
            other,
            policy,
            entry.get("classification", "missing_evidence"),
            entry.get("evidence", ""),
        )
        results.append(
            {
                "project_id": key[0],
                "component": key[1],
                "workbook": asdict(reference),
                "comparison": asdict(other),
                "basis_evidence": entry["basis_evidence"],
                **asdict(result),
            }
        )
    # Canonical monetary strings; no float JSON conversion.
    return json.loads(json.dumps(results, default=str))


def build(root: Path, output: Path, comparison_path: Path | None = None) -> dict:
    from openpyxl import load_workbook

    rows, inventory, monthly = workbook_rows(root)
    inputs = {str(CONTRACT): digest(root / CONTRACT)}
    inputs.update(
        {
            str(FDAS / name): profile["sha256"]
            for name, profile in inventory.items()
            if isinstance(profile, dict) and "sha256" in profile
        }
    )
    crosswalk_path = CURATED / "fdas_project_cost_crosswalk.csv"
    inputs[str(crosswalk_path)] = digest(root / crosswalk_path)
    with (root / crosswalk_path).open(newline="", encoding="utf-8") as stream:
        crosswalk = list(csv.DictReader(stream))
    requirement_path = CURATED / "cost_requirement_identity.csv"
    inputs[str(requirement_path)] = digest(root / requirement_path)
    with (root / requirement_path).open(newline="", encoding="utf-8") as stream:
        requirements = {r["OPAQUE_ID"] for r in csv.DictReader(stream)}
    links = []
    for index, row in enumerate(crosswalk, 2):
        if row["WORKBOOK_SHA256"] != inventory[row["WORKBOOK_FILE"]]["sha256"]:
            raise ValueError("existing cell crosswalk digest mismatch")
        if not set(filter(None, row["REQUIREMENT_IDS"].split("|"))) <= requirements:
            raise ValueError("existing cell crosswalk has an unknown requirement")
        if row["WORKBOOK_VALUE"]:
            book = load_workbook(
                root / FDAS / row["WORKBOOK_FILE"], read_only=True, data_only=True
            )
            try:
                value = book[row["WORKBOOK_SHEET"]][row["WORKBOOK_CELL"]].value
                if _number(value) != Decimal(row["WORKBOOK_VALUE"]):
                    raise ValueError("existing cell crosswalk value mismatch")
            finally:
                book.close()
        links.append(
            {
                "project_id": row["PROJECT_ID"],
                "requirement_ids": row["REQUIREMENT_IDS"],
                "source": f"fdas_project_cost_crosswalk.csv:row-{index}",
                "workbook": row["WORKBOOK_FILE"],
                "cell": row["WORKBOOK_CELL"],
                "mapping_status": row["MAPPING_STATUS"],
                "unit": row["VALUE_UNIT"],
                "provenance": "workbook_assumption",
            }
        )
    inventory["requirement_links"] = links
    assets = []
    cost_map = Path("reports/cost/big_foot_cost_map_reconciliation.csv")
    if (root / cost_map).is_file():
        inputs[str(cost_map)] = digest(root / cost_map)
        with (root / cost_map).open(newline="", encoding="utf-8") as stream:
            # No private locators, workbook core properties or project names are published.
            fields = [
                "project_id",
                "direction",
                "requirement_id",
                "award_id",
                "row_kind",
                "value_low_mm",
                "value_high_mm",
                "evidence_derivation",
                "counting_disposition",
            ]
            assets = [
                {field: row[field] for field in fields}
                for row in csv.DictReader(stream)
                if row["row_kind"] != "fdas_crosswalk"
            ]
    upstream = []
    for path in UPSTREAM:
        sha = digest(root / path) if (root / path).is_file() else ""
        if sha:
            inputs[path] = sha
        upstream.append(
            {
                "path": path,
                "sha256": sha,
                "status": "NOT_EVALUATED",
                "reason": (
                    "contract consumer not implemented"
                    if sha
                    else "required manifest unavailable"
                ),
            }
        )
    annual = defaultdict(lambda: {component: [] for component in MONTHLY.values()})
    for row in monthly:
        for component in MONTHLY.values():
            annual[(row["project_id"], row["period"][:4])][component].append(
                Decimal(row[component])
            )
    traces = [
        {
            "project_id": key[0],
            "period": key[1],
            **{
                component: str(_sum(values)) for component, values in components.items()
            },
        }
        for key, components in sorted(annual.items())
    ]
    comparisons = []
    if comparison_path is not None:
        comparisons = compare_rows(rows, json.loads(comparison_path.read_text()))
        inputs["external_comparisons.json"] = digest(comparison_path)
    payload = {
        "rows": rows,
        "inventory": inventory,
        "assets": assets,
        "traces": traces,
        "comparisons": comparisons,
        "upstream": upstream,
    }
    artifacts = {
        "reconciliation.csv": _csv(rows, FIELDS),
        "monthly.csv": _csv(monthly, ["project_id", "period", *MONTHLY.values()]),
        "reconciliation.json": canonical(payload),
        "index.html": render_html(payload),
    }
    # Freeze implementation as well as evidence inputs; no machine path or wall clock.
    for path in sorted(Path(__file__).parent.glob("*.py")):
        inputs[path.relative_to(root).as_posix()] = digest(path)
    builder = Path("scripts/cost/build_fdas_workbook_reconciliation.py")
    inputs[str(builder)] = digest(root / builder)
    manifest = {
        "contract_version": "1.0",
        "generated_at_policy": "omitted for deterministic output",
        "decimal_policy": "native Decimal strings; exact sums; one decimal in HTML only",
        "materiality_policy": (
            json.loads(comparison_path.read_text()).get("policy")
            if comparison_path
            else None
        ),
        "upstream": upstream,
        "inputs": dict(sorted(inputs.items())),
        "outputs": {
            name: hashlib.sha256(content.encode()).hexdigest()
            for name, content in artifacts.items()
        },
        "known_limits": [
            "No full V30 financial-engine reproduction",
            "V50 adoption excluded (#899)",
            "External project identity, scope and basis equivalence not established",
            "Stakeholder circulation requires maintainer approval (#1017)",
        ],
    }
    output.mkdir(parents=True, exist_ok=True)
    # Refuse any destination that could overwrite a source workbook or input.
    if any(
        (output / name).resolve() == (root / source).resolve()
        for name in [*artifacts, "manifest.json"]
        for source in inputs
    ):
        raise ValueError("output destination overlaps a frozen input")
    for name, content in {**artifacts, "manifest.json": canonical(manifest)}.items():
        (output / name).write_text(content, encoding="utf-8")
    return manifest
