"""Read-only, digest-pinned workbook inventory and V30 accounting bridges."""

import hashlib
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook


def digest(data):
    return hashlib.sha256(data).hexdigest()


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    if value is None:
        return None
    return str(value)


def read_cell(path, sheet, coordinate):
    """Expose formula/cache independently; never evaluate a formula.

    Source: openpyxl 3.1 reader.excel.load_workbook, data_only parameter.
    https://openpyxl.readthedocs.io/en/stable/api/openpyxl.reader.excel.html
    """
    data = path.read_bytes()
    formula_book = load_workbook(BytesIO(data), read_only=True, data_only=False)
    cache_book = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        cell = formula_book[sheet][coordinate]
        formula = cell.value if cell.data_type == "f" else None
        value = cache_book[sheet][coordinate].value
        return {
            "formula": formula,
            "value": scalar(value),
            "status": (
                "NOT_EVALUATED"
                if value is None
                else "CACHED_UNVERIFIED" if formula else "READ"
            ),
        }
    finally:
        formula_book.close()
        cache_book.close()


def freeze_inventory(root, contract):
    """Verify both bytes and schema against an explicitly frozen contract.

    Core document metadata is deliberately not read or exported.
    """
    inventory = {"vintage": contract["vintage"], "workbooks": {}}
    for name, expected in sorted(contract["workbooks"].items()):
        if name not in {"lease_assumptions.xlsx", "financial_project_summary.xlsx"}:
            raise ValueError("unexpected workbook name")
        data = (root / name).read_bytes()
        if digest(data) != expected["sha256"]:
            raise ValueError(f"digest mismatch: {name}")
        book = load_workbook(BytesIO(data), read_only=True, data_only=False)
        try:
            sheets = {
                s.title: {
                    "rows": s.max_row,
                    "columns": s.max_column,
                    "headers": [c.value for c in next(s.iter_rows())],
                }
                for s in book
            }
            if sheets != expected["sheets"]:
                raise ValueError(f"schema drift: {name}")
            inventory["workbooks"][name] = {"sha256": digest(data), "sheets": sheets}
        finally:
            book.close()
    return inventory


# Source: frozen financial_project_summary.xlsx, Project_Summary row 1.
# Subtotals are non-additive; OPEX and financial results never enter CAPEX.
SUMMARY_MAP = {
    "E": ("host", "included"),
    "F": ("SURF", "included"),
    "G": ("pumps", "included"),
    "H": ("pumps", "included"),
    "I": ("dry-system", "included"),
    "J": ("facilities", "subtotal"),
    "K": ("drilling", "included"),
    "L": ("completion", "included"),
    "M": ("D&C", "subtotal"),
    "N": ("other", "excluded"),
    "O": ("other", "excluded"),
    "P": ("OPEX", "excluded"),
    "Q": ("OPEX", "excluded"),
    "R": ("other", "excluded"),
    "S": ("other", "excluded"),
}
MONTHLY_MAP = {
    "facilities": ("Facilities_USD", "J"),
    "drilling": ("DRILLING_COST_USD", "K"),
    "completion": ("COMPLETION_COST_USD", "L"),
    "variable_opex": ("Variable_OPEX_USD", "P"),
    "fixed_opex": ("Fixed_OPEX_USD", "Q"),
}


def _number(cell):
    if cell.value is None or cell.data_type in {"f", "e", "b"}:
        return None
    value = Decimal(str(cell.value))
    if not value.is_finite():
        raise ValueError("nonfinite workbook value")
    return value


def _sum(values):
    values = list(values)
    return None if any(v is None for v in values) else sum(values, Decimal(0))


def _difference(left, right):
    return None if left is None or right is None else scalar(left - right)


def _monthly(sheet, summary, row):
    records = list(sheet.iter_rows())
    headers = [c.value for c in records[0]]
    months = []
    for cells in records[1:]:
        if not isinstance(cells[0].value, (date, datetime)):
            raise ValueError("monthly trace requires dated rows")
        months.append(
            {
                "period": scalar(cells[0].value)[:7],
                **{
                    category: scalar(_number(cells[headers.index(header)]))
                    for category, (header, _) in MONTHLY_MAP.items()
                },
            }
        )
    periods = [m["period"] for m in months]
    if periods != sorted(set(periods)):
        raise ValueError("monthly periods must be ordered and unique")
    checks = {}
    for category, (_, column) in MONTHLY_MAP.items():
        total = _sum(
            Decimal(m[category]) if m[category] is not None else None for m in months
        )
        variance = _difference(total, _number(summary[f"{column}{row}"]))
        checks[category] = {
            "monthly_sum_usd": scalar(total),
            "variance_usd": variance,
            "status": (
                "EXACT"
                if variance == "0"
                else "NOT_EVALUATED" if variance is None else "UNRESOLVED"
            ),
        }
    return months, checks


def _project(summary, row, assumptions, project_sheet):
    project, system = summary[f"A{row}"].value, summary[f"B{row}"].value
    months, checks = _monthly(project_sheet, summary, row)
    columns = {c.value: c.column_letter for c in next(assumptions.iter_rows())}
    if system not in columns:
        raise ValueError("unmapped development system")
    # Source: V30 generator, facilities booking block: injection facility is
    # booked once at first oil, absent for non-producing projects. SI mega=10^6;
    # workbook assumptions row 15 explicitly specifies millions of USD.
    injection = (
        _number(assumptions[f"{columns[system]}15"]) * Decimal("1000000")
        if summary[f"C{row}"].value is not None
        else Decimal(0)
    )
    visible = _sum(_number(summary[f"{c}{row}"]) for c in ("E", "F", "G", "H", "I"))
    facilities = _number(summary[f"J{row}"])
    drilling, completion = _number(summary[f"K{row}"]), _number(summary[f"L{row}"])
    crosswalk = [
        {
            "cell": f"Project_Summary!{column}{row}",
            "parameter": summary[f"{column}1"].value,
            "category": category,
            "disposition": disposition,
            "unit": "USD",
            "disclosed": None,
            "mapped": None,
            "modeled": None,
            "variance": None,
            "comparison_status": "NOT_EVALUATED",
            "requirement_ids": [],
            "mapping_status": "unmapped",
            "value": scalar(_number(summary[f"{column}{row}"])),
        }
        for column, (category, disposition) in SUMMARY_MAP.items()
    ]
    return {
        "project": project,
        "system": system,
        "vintage": "FDAS_V30",
        "provenance": "workbook_assumption",
        "unit": "USD",
        "currency": "USD",
        "basis_year": None,
        "ownership": None,
        "time_basis": f"{months[0]['period']}/{months[-1]['period']}",
        "identity_status": "workbook identity only; external scope crosswalk required",
        "comparison_status": "NOT_EVALUATED",
        "comparison_reason": "independently evidenced identity, scope, price/ownership basis and materiality policy required",
        "water_injection_facility_usd": scalar(injection),
        "injection_source": f"lease_assumptions.xlsx:assumptions!{columns[system]}15",
        "component_bridge_variance_usd": _difference(
            facilities, _sum((visible, injection))
        ),
        "dnc_bridge_variance_usd": _difference(
            _number(summary[f"M{row}"]), _sum((drilling, completion))
        ),
        "total_capex_usd": scalar(_sum((facilities, drilling, completion))),
        "crosswalk": crosswalk,
        "months": months,
        "monthly_checks": checks,
    }


def reconcile_workbooks(root, contract):
    inventory = freeze_inventory(root, contract)
    if contract["vintage"] != "FDAS_V30":
        raise ValueError("V50 adoption is outside this reconciliation")
    snapshots = {name: (root / name).read_bytes() for name in inventory["workbooks"]}
    if any(
        digest(data) != inventory["workbooks"][name]["sha256"]
        for name, data in snapshots.items()
    ):
        raise ValueError("workbook digest mismatch before parsing")
    books = {
        name: load_workbook(BytesIO(data), read_only=True, data_only=False)
        for name, data in snapshots.items()
    }
    try:
        summary = books["financial_project_summary.xlsx"]["Project_Summary"]
        assumptions = books["lease_assumptions.xlsx"]["assumptions"]
        projects = [
            _project(
                summary,
                row,
                assumptions,
                books["financial_project_summary.xlsx"][summary[f"A{row}"].value],
            )
            for row in range(2, summary.max_row + 1)
        ]
        assumption_rows = list(assumptions.iter_rows())
        parameters = [
            {
                "parameter": row[0].value,
                "system": assumption_rows[0][column].value,
                "cell": f"{assumption_rows[0][column].column_letter}{row_number}",
                "value": scalar(cell.value),
                "unit_note": row[-1].value,
                "provenance": "workbook_assumption",
                "vintage": contract["vintage"],
            }
            for row_number, row in enumerate(assumption_rows[1:], start=2)
            for column, cell in enumerate(row[1:-1], start=1)
        ]
    finally:
        for book in books.values():
            book.close()
    # A changed file during ingestion invalidates the entire result.
    if freeze_inventory(root, contract) != inventory:
        raise ValueError("workbook changed during ingestion")
    return {
        "inventory": inventory,
        "parameters": parameters,
        "projects": sorted(projects, key=lambda p: p["project"]),
        "comparisons": [],
        "upstream": {},
        "v30_status": "workbook component/monthly consistency only; full financial reproduction NOT_EVALUATED",
        "v50_status": "NOT_EVALUATED; changes remain under issue 899",
    }
