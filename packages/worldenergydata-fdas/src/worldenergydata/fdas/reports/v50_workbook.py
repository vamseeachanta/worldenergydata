"""Export V50 source objects without selecting unresolved economic assumptions.

Schema authority: docs/modules/bsee/analysis/production/FDAS_V30/
financial_project_summary.xlsx, Project_Summary header row. Required provenance
and separate tax bases follow repository issue 901. No economics are calculated.
"""

from dataclasses import dataclass
from datetime import date, datetime
from math import isfinite
from numbers import Real
from pathlib import Path
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


# Frozen lineage schema, including its original units and column order.
LEGACY_SUMMARY_COLUMNS = (
    "Project Name",
    "DEV SYSTEM",
    "FO Month",
    "TOTAL OIL BBL",
    "Host CAPEX USD",
    "SURF USD",
    "Booster Pump USD",
    "Water Injection Pump USD",
    "Dry Well System USD",
    "Facilities Cost USD",
    "DRILLING COST USD",
    "COMPLETION COST USD",
    "Drilling and Completion Total USD",
    "Revenue USD",
    "Royalty USD",
    "Variable OPEX USD",
    "Fixed OPEX USD",
    "Net CashFlow USD",
    "NPV USD",
    "MIRR monthly",
    "MIRR annual",
    "PRODUCER WELLS",
    "INJECTOR WELLS",
    "Booster Pump Count",
    "Water Injection Pump Count",
    "TOTAL WELLBORES",
    "TOTAL DRILLING DAYS",
    "TOTAL COMPLETION DAYS",
)
SUMMARY_COLUMNS = (*LEGACY_SUMMARY_COLUMNS, "After-tax NPV USD", "NPV Reference Date")
_TEXT_COLUMNS = {"Project Name", "DEV SYSTEM"}
_DATE_COLUMNS = {"FO Month", "NPV Reference Date"}
_NULLABLE_COLUMNS = {"FO Month", "MIRR monthly", "MIRR annual"}


@dataclass(frozen=True)
class NotApplicable:
    """An intentionally empty first-oil or MIRR cell with a required explanation."""

    reason: str

    def __post_init__(self):
        if not isinstance(self.reason, str) or not self.reason.strip():
            raise ValueError("Not-applicable reason must be nonblank text")


@dataclass(frozen=True)
class V50WorkbookMetadata:
    """Required caller-supplied provenance; no economic defaults are provided."""

    assumption_version: str
    git_commit: str
    production_cutoff: date
    wti_deck_source: str
    npv_reference_convention: str
    generated_at: datetime

    def as_rows(self):
        """Validate provenance and return the public workbook metadata rows."""
        for name in (
            "assumption_version",
            "wti_deck_source",
            "npv_reference_convention",
        ):
            value = getattr(self, name)
            if not isinstance(value, str) or not value.strip():
                raise ValueError(f"{name} must be nonblank text")
        # Full Git object names; abbreviated or unknown revisions are ambiguous.
        if not isinstance(self.git_commit, str) or not re.fullmatch(
            r"(?:[0-9a-f]{40}|[0-9a-f]{64})", self.git_commit
        ):
            raise ValueError("git_commit must be a full hexadecimal Git object name")
        if type(self.production_cutoff) is not date:
            raise ValueError("production_cutoff must be a date")
        if (
            not isinstance(self.generated_at, datetime)
            or self.generated_at.utcoffset() is None
        ):
            raise ValueError("generated_at must be a timezone-aware datetime")
        return {
            "assumption_version": self.assumption_version,
            "git_commit": self.git_commit,
            "production_cutoff": self.production_cutoff.isoformat(),
            "wti_deck_source": self.wti_deck_source,
            "npv_reference_convention": self.npv_reference_convention,
            "generated_at": self.generated_at.isoformat(),
            "pre_tax_npv_column": "NPV USD",
            "after_tax_npv_column": "After-tax NPV USD",
            "tax_basis:Net CashFlow USD": "pre-tax",
            "tax_basis:MIRR monthly": "pre-tax",
            "tax_basis:MIRR annual": "pre-tax",
            "npv_units": "USD",
            "canonical_acceptance_status": "NOT_EVALUATED",
            "acceptance_limitation": "Assumptions and NPV convention require resolution of issues 899 and 900",
            "lineage": "V30 is lineage-only",
        }


def _validated_rows(summary):
    """Validate values and collect reasons for explicitly not-applicable cells."""
    if not summary.columns.is_unique:
        raise ValueError("Summary columns must be unique")
    missing = set(SUMMARY_COLUMNS) - set(summary.columns)
    if missing:
        raise ValueError(f"Missing required columns: {sorted(missing)}")
    if set(summary.columns) != set(SUMMARY_COLUMNS):
        raise ValueError("Unexpected summary columns; map source objects explicitly")
    if summary.empty:
        raise ValueError("Summary must not be empty")
    if summary["Project Name"].duplicated().any():
        raise ValueError("Project names must be unique")
    rows, reasons = [], {}
    source_rows = summary.loc[:, SUMMARY_COLUMNS].itertuples(index=False, name=None)
    for row_number, row in enumerate(source_rows, start=2):
        values = []
        for column_number, (name, value) in enumerate(
            zip(SUMMARY_COLUMNS, row), start=1
        ):
            if isinstance(value, NotApplicable):
                if name not in _NULLABLE_COLUMNS:
                    raise ValueError(f"{name} is not nullable")
                cell = f"{get_column_letter(column_number)}{row_number}"
                reasons[f"not_applicable:Project_Summary!{cell}"] = value.reason
                values.append(None)
                continue
            if name in _TEXT_COLUMNS:
                if not isinstance(value, str) or not value.strip():
                    raise ValueError(f"{name} must be nonblank text")
            elif name in _DATE_COLUMNS:
                if not isinstance(value, date) or pd.isna(value):
                    raise ValueError(f"{name} must be a date")
                if isinstance(value, datetime) and value.tzinfo is not None:
                    raise ValueError(f"{name} must not contain a timezone")
            elif (
                isinstance(value, bool)
                or not isinstance(value, Real)
                or not isfinite(value)
            ):
                raise ValueError(f"{name} must be a finite numeric value")
            values.append(value)
        rows.append(values)
    return rows, reasons


def _append_literal(sheet, values):
    """Keep source text literal so spreadsheet formula prefixes cannot execute."""
    sheet.append(list(values))
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def write_v50_workbook(
    output_path: str | Path, summary: pd.DataFrame, metadata: V50WorkbookMetadata
) -> Path:
    """Write a candidate Project_Summary and provenance sheet, without overwrite.

    The input table uses the frozen lineage headers plus After-tax NPV USD and
    NPV Reference Date. Both NPVs are required engine outputs in USD, already
    discounted on the supplied convention. Net cashflow and both MIRRs are
    pre-tax. Only first-oil and MIRR cells permit NotApplicable(reason), written
    as blanks with per-cell explanations in Metadata. This export does not
    establish an accepted canonical model or supply missing fiscal terms. Until
    the basis decisions are resolved, acceptance remains NOT_EVALUATED in every export.
    """
    info = metadata.as_rows()
    rows, reasons = _validated_rows(summary)
    info.update(reasons)
    book = Workbook()
    sheet = book.active
    sheet.title = "Project_Summary"
    _append_literal(sheet, SUMMARY_COLUMNS)
    for row in rows:
        _append_literal(sheet, row)
    provenance = book.create_sheet("Metadata")
    for row in info.items():
        _append_literal(provenance, row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    path = Path(output_path)
    # Exclusive creation preserves previously generated evidence.
    with path.open("xb") as stream:
        book.save(stream)
    book.close()
    return path
