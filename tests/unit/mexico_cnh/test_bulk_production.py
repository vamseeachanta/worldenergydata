"""Bulk ingest regression tests; workbook layout follows SENER January 2025.

Source: https://hidrocarburos.energia.gob.mx/estadisticas/ (2025 workbook),
Aceite y condensado and Gas Natural, units row and header row.
All rates below are synthetic inputs; expected volumes are hand calculations.
"""

from datetime import datetime

import pytest
from openpyxl import Workbook

from worldenergydata.mexico_cnh.bulk import load_production_workbook


def workbook(path, *, oil=2, gas=3, condensate=1):
    """Synthetic February 2024: oil=2 Mbd, gas=3 MMpcd, condensate=1 Mbd."""
    book = Workbook()
    liquid = book.active
    liquid.title = "Aceite y condensado"
    liquid.append(["Producción de petróleo y condensado por campo"])
    liquid.append(["Unidades: miles de barriles diarios (Mbd)"])
    liquid.append(
        [
            "Fecha",
            "Cuenca",
            "Ubicación",
            "Campo",
            "Operador",
            "Líquidos (Mbd)",
            "Petróleo (Mbd)",
            "Condensado (Mbd)",
        ]
    )
    liquid.append(
        [
            datetime(2024, 2, 1),
            "Basin",
            "Terrestre",
            "EXAMPLE",
            "OP",
            3,
            oil,
            condensate,
        ]
    )
    natural = book.create_sheet("Gas Natural")
    natural.append(["Unidades: millones de pies cúbicos por día (MMpcd)"])
    natural.append(
        [
            "Fecha",
            "Cuenca",
            "Ubicación",
            "Campo",
            "Operador",
            "Gas natural sin nitrógeno (MMpcd)",
            "Nitrógeno (MMpcd)",
        ]
    )
    natural.append(
        [datetime(2024, 2, 1), "Basin", "Terrestre", "EXAMPLE", "OP", gas, None]
    )
    book.save(path)
    book.close()
    return path


def test_monthly_volumes_and_provenance(tmp_path):
    """Leap February: 2*1000*29=58000 bbl; 3*1000*29=87000 Mcf.

    Condensate: 1*1000*29=29000 bbl, kept separate from petroleum.
    Factors follow the cited workbook's Mbd and MMpcd units, not gas BOE.
    """
    result = load_production_workbook(workbook(tmp_path / "input.xlsx"))
    row = result.production.iloc[0]
    assert row.oil_bbl == 58000
    assert row.gas_mcf == 87000
    assert row.condensate_bbl == 29000
    assert row.month == 2 and row.year == 2024
    assert row.field_name == "EXAMPLE"
    assert len(result.sha256) == 64
    assert result.vintage == "2024-02"
    assert row.water_bbl != row.water_bbl  # unknown is NaN, never zero


@pytest.mark.parametrize("rate", [-1, "Infinity", "not a measurement", True])
def test_invalid_measurements_fail_closed(tmp_path, rate):
    with pytest.raises((TypeError, ValueError), match="rates"):
        load_production_workbook(workbook(tmp_path / "input.xlsx", oil=rate))


def test_blank_condensate_stays_unknown(tmp_path):
    result = load_production_workbook(
        workbook(tmp_path / "input.xlsx", condensate=None)
    )
    assert result.production.condensate_bbl.isna().all()
    assert result.production.iloc[0].oil_bbl == 58000


@pytest.mark.parametrize("change", ["unit", "duplicate", "date", "identity", "sheet"])
def test_schema_and_identity_drift_fail_closed(tmp_path, change):
    from openpyxl import load_workbook

    path = workbook(tmp_path / "input.xlsx")
    book = load_workbook(path)
    sheet = book["Aceite y condensado"]
    if change == "unit":
        sheet.cell(3, 7).value = "Petróleo (bpd)"
    elif change == "duplicate":
        sheet.append([cell.value for cell in sheet[4]])
    elif change == "date":
        sheet.cell(4, 1).value = datetime(2024, 2, 2)
    elif change == "identity":
        sheet.cell(4, 4).value = None
    else:
        book["Gas Natural"].title = "Unknown"
    book.save(path)
    book.close()
    with pytest.raises(ValueError):
        load_production_workbook(path)


def test_unmatched_gas_identity_preserves_partial_records(tmp_path):
    from openpyxl import load_workbook

    path = workbook(tmp_path / "input.xlsx")
    book = load_workbook(path)
    book["Gas Natural"].cell(3, 4).value = "OTHER"
    book.save(path)
    book.close()
    result = load_production_workbook(path).production
    assert len(result) == 2
    oil = result[result.field_name == "EXAMPLE"].iloc[0]
    assert oil.oil_bbl == 58000 and oil.gas_mcf != oil.gas_mcf


def test_overflowed_monthly_volume_is_rejected(tmp_path):
    """A finite synthetic input can overflow after unit and calendar scaling."""
    with pytest.raises(ValueError, match="finite"):
        load_production_workbook(workbook(tmp_path / "input.xlsx", oil=1e308))
