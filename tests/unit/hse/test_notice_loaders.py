"""Offline contracts for public notice intake; every entity is synthetic.

Field labels: OPRED PON1 notifications 4966 and 6788 (2024), read 2026-10-22.
The seed quantity, cause and date are also in issue #1325 (2026), Seed event.
HSE workbook headers below are an explicit test mapping, not a claim about the
unread spreadsheet. Expected dates/quantities are transcribed fixture literals;
no expected result is obtained from the implementation under test.
"""

import json
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from worldenergydata.hse.importers.notice_loaders import (
    HSE_SPREADSHEET_URL,
    load_hse_notices,
    manual_hse_step,
    parse_pon1,
)
from worldenergydata.hse.importers.offshore_importer import OffshoreImporter
from worldenergydata.hse.offshore import Event, SourceDocument

FIXTURES = Path(__file__).parent / "fixtures" / "notices"
ACCESS = date(2026, 10, 22)
PON1_URL = "https://itportal.energysecurity.gov.uk/irs/publications/pon1/2024/notification/4966"
COLUMNS = {
    "source_ref": "Reference",
    "title": "Notice title",
    "doc_type": "Type",
    "issue_date": "Issued",
    "compliance_date": "Compliance",
    "status": "Status",
    "legal_provision": "Provision",
}


@pytest.fixture
def reviewed():
    # Hand-curated synthetic inventory, never derived from incoming raw text.
    return json.loads((FIXTURES / "reviewed_text.json").read_text())


@pytest.fixture
def context():
    return {
        "event_id": "synthetic-seed",
        "country": "GB",
        "basin_or_area": "synthetic area",
        "installation_aliases": ["synthetic-alias"],
        "operation": "riser or BOP handling",
        "equipment": ["riser"],
        "failure_mode": "loss of containment",
        "failure_summary": "The riser ruptured and released chemical.",
    }


@pytest.fixture
def html():
    return (FIXTURES / "pon1_synthetic.html").read_text()


def parse(html, context, reviewed):
    return parse_pon1(
        html, context=context, url=PON1_URL, access_date=ACCESS, reviewed_text=reviewed
    )


def workbook(tmp_path, rows=None, headers=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "notices"
    sheet.append(headers or list(COLUMNS.values()))
    for row in rows or [
        [
            "synthetic-notice",
            "synthetic notice",
            "improvement notice",
            datetime(2024, 3, 21),
            datetime(2024, 4, 21),
            "complied with",
            "synthetic provision",
        ]
    ]:
        sheet.append(row)
    path = tmp_path / "synthetic.xlsx"
    book.save(path)
    book.close()
    return path


def load(path, reviewed, **changes):
    options = {
        "columns": COLUMNS,
        "sheet": "notices",
        "header_row": 1,
        "associations": {"synthetic-notice": "synthetic-seed"},
        "access_date": ACCESS,
        "reviewed_text": reviewed,
        "date_format": "%d/%m/%Y",
    }
    options.update(changes)
    return load_hse_notices(path, **options)


def test_seed_chemical_preserves_units_date_and_unknowns(html, context, reviewed):
    """Literal 93.87 kg, 2024-02-01 and cause: issue #1325, Seed event.

    First observed time has no zone; hand expectation is a date, not a UTC time.
    No injury/fatality count or publication date is published in this layout.
    """
    packet = parse(html, context, reviewed)
    event = Event.model_validate(packet["events"][0])
    source = SourceDocument.model_validate(packet["sources"][0])
    assert event.release_quantity == 93.87
    assert event.release_unit == "kg"
    assert event.event_date == date(2024, 2, 1)
    assert event.causes_as_stated == ["Riser - Total Rupture"]
    assert event.block_or_lease == "204/20"
    assert event.injuries is event.fatalities is None
    assert source.issue_date is None
    assert source.source_ref == "IRS/2024/4023/PON1"
    assert "06:45" in " ".join(packet["provenance_notes"])


def test_oil_layout_keeps_published_quantity_without_conversion(
    html, context, reviewed
):
    """Oil labels/unit: OPRED notification 6788 (2024), read 2026-10-22.

    Synthetic 0.5 tonnes stays 0.5 tonnes: no density or unit factor is assumed.
    """
    oil = (
        html.replace("Chemical release", "Oil release")
        .replace("Chemical product released", "Oil product released")
        .replace(
            "Total quantity of chemical released", "Total quantity of oil released"
        )
    )
    oil = oil.replace("93.87 kg", "0.5 tonnes")
    reviewed += ["Oil release", "0.5 tonnes", "tonnes"]
    row = parse(oil, context, reviewed)["events"][0]
    assert row["release_quantity"] == 0.5 and row["release_unit"] == "tonnes"


@pytest.mark.parametrize("quantity", ["-1 kg", "NaN kg", "93.87", "93.87 kg extra"])
def test_malformed_release_is_not_silently_coerced(html, context, reviewed, quantity):
    with pytest.raises(ValueError, match="PON1"):
        parse(html.replace("93.87 kg", quantity), context, reviewed + [quantity])


def test_repeated_chemical_is_refused_instead_of_losing_or_summing_it(
    html, context, reviewed
):
    extra = "<tr><th>Chemical product released</th><td>synthetic chemical</td></tr>"
    with pytest.raises(ValueError, match="PON1"):
        parse(
            html.replace("</main>", f"<table>{extra}</table></main>"), context, reviewed
        )


@pytest.mark.parametrize(
    "old,new",
    [
        ("First observed date", "Changed label"),
        ("Open Government Licence v3.0", "licence not established"),
        ("IRS/2024/4023/PON1", "unverified reference"),
        ("01 February 2024", "not a date"),
    ],
)
def test_missing_schema_or_licence_fails_closed(html, context, reviewed, old, new):
    with pytest.raises(ValueError):
        parse(html.replace(old, new), context, reviewed + [new])


def test_context_cannot_overwrite_published_release(html, context, reviewed):
    with pytest.raises(ValueError, match="context"):
        parse(html, dict(context, release_quantity=0), reviewed)


@pytest.mark.parametrize(
    "field", ["Installation", "Field", "Source", "Organisations involved"]
)
def test_personal_name_in_raw_free_text_fails_without_echo(
    html, context, reviewed, field
):
    # Invented tokens are assembled only at runtime; no person's name is stored.
    injected = "Synthetic Individual"
    import re

    altered = re.sub(
        rf"(<th[^>]*>{field}</th><td>).*?(</td>)", rf"\g<1>{injected}\2", html
    )
    with pytest.raises(ValueError, match="privacy") as caught:
        parse(altered, context, reviewed)
    assert injected not in str(caught.value)


def test_citation_has_all_components_and_exact_attribution(html, context, reviewed):
    source = SourceDocument.model_validate(parse(html, context, reviewed)["sources"][0])
    assert source.citation == (
        'OPRED, release notice IRS/2024/4023/PON1, "IRS/2024/4023/PON1", '
        f"issue date not published. {PON1_URL} (accessed 2026-10-22). "
        "Licence: OGL-3.0. Contains public sector information licensed under "
        "the Open Government Licence v3.0."
    )


def test_spreadsheet_mapping_and_dates_do_not_create_an_event(tmp_path, reviewed):
    result = load(workbook(tmp_path), reviewed)
    assert result.status == "ready" and not result.manual_urls
    source = result.sources[0]
    assert source.event_id == "synthetic-seed"
    assert source.issue_date == date(2024, 3, 21)
    assert source.compliance_date == date(2024, 4, 21)
    assert source.doc_type == "improvement notice"
    assert source.url == HSE_SPREADSHEET_URL
    assert source.legal_provision == "synthetic provision"
    assert "not the event date" in source.association_basis


def test_prohibition_notice_and_text_dates(tmp_path, reviewed):
    rows = [
        [
            "synthetic-notice",
            "synthetic notice",
            "prohibition notice",
            "21/03/2024",
            None,
            "complied with",
            None,
        ]
    ]
    result = load(
        workbook(tmp_path, rows), reviewed + ["prohibition notice", "21/03/2024"]
    )
    assert result.sources[0].doc_type == "prohibition notice"
    assert result.sources[0].issue_date == date(2024, 3, 21)
    assert result.sources[0].compliance_date is None


@pytest.mark.parametrize("kind", ["missing", "invalid", "wrong_extension"])
def test_unreadable_workbook_outputs_manual_urls(tmp_path, reviewed, kind):
    path = tmp_path / ("input.xls" if kind == "wrong_extension" else "input.xlsx")
    if kind != "missing":
        path.write_bytes(b"not a workbook")
    result = load(path, reviewed)
    assert result.status == "manual_required" and not result.sources
    assert HSE_SPREADSHEET_URL in result.manual_urls
    assert "https://notices.hse.gov.uk/" in result.manual_urls


def test_blocked_live_spreadsheet_is_a_manual_step():
    result = manual_hse_step("BLOCKED-NETWORK")
    assert result.reason == "BLOCKED-NETWORK" and not result.sources
    assert len(result.manual_urls) == len(set(result.manual_urls))


@pytest.mark.parametrize(
    "change", ["missing_header", "duplicate_header", "missing_sheet"]
)
def test_unknown_workbook_layout_requires_review(tmp_path, reviewed, change):
    headers = list(COLUMNS.values())
    if change == "missing_header":
        headers[0] = "Changed reference"
    if change == "duplicate_header":
        headers[-1] = headers[0]
    path = workbook(tmp_path, headers=headers)
    with pytest.raises(ValueError, match="layout"):
        load(
            path, reviewed, **({"sheet": "absent"} if change == "missing_sheet" else {})
        )


def test_missing_selected_notice_does_not_claim_success(tmp_path, reviewed):
    with pytest.raises(ValueError, match="association"):
        load(workbook(tmp_path), reviewed, associations={"missing": "synthetic-seed"})


def test_personal_name_in_notice_title_fails(tmp_path, reviewed):
    rows = [
        [
            "synthetic-notice",
            "Synthetic Individual",
            "improvement notice",
            date(2024, 3, 21),
            None,
            "complied with",
            None,
        ]
    ]
    with pytest.raises(ValueError, match="privacy"):
        load(workbook(tmp_path, rows), reviewed)


def test_person_columns_are_refused_even_if_text_was_approved(tmp_path, reviewed):
    rows = [
        [
            "synthetic-notice",
            "synthetic notice",
            "improvement notice",
            date(2024, 3, 21),
            None,
            "complied with",
            None,
            "synthetic inspector",
        ]
    ]
    with pytest.raises(ValueError, match="privacy"):
        load(
            workbook(tmp_path, rows, [*COLUMNS.values(), "Inspector"]),
            reviewed + ["synthetic inspector"],
        )


def test_workbook_formula_is_not_evaluated_or_admitted(tmp_path, reviewed):
    rows = [
        [
            "synthetic-notice",
            '=CONCAT("synthetic", " notice")',
            "improvement notice",
            date(2024, 3, 21),
            None,
            "complied with",
            None,
        ]
    ]
    with pytest.raises(ValueError, match="formula"):
        load(workbook(tmp_path, rows), reviewed)


def test_loader_duplicate_notice_reference_is_ambiguous(tmp_path, reviewed):
    row = [
        "synthetic-notice",
        "synthetic notice",
        "improvement notice",
        date(2024, 3, 21),
        None,
        "complied with",
        None,
    ]
    with pytest.raises(ValueError, match="duplicate"):
        load(workbook(tmp_path, [row, row]), reviewed)


def test_unique_key_rerun_and_planted_alias_duplicate_export_log(
    tmp_path, html, context, reviewed
):
    """Rule: #1325 (2026), Deduplication: same country/name, date ±2 days.

    Independent hand check: February 3 minus February 1 = two days. A matching
    alias and GB therefore give one event, two distinct source references.
    Rerunning either source adds zero rows by (source, source_ref) uniqueness.
    """
    importer = OffshoreImporter(tmp_path / "offshore.sqlite")
    packet = parse(html, context, reviewed)
    assert importer.ingest(packet)["sources_added"] == 1
    assert importer.ingest(packet)["sources_added"] == 0
    second = json.loads(json.dumps(packet))
    second["events"][0].update(
        event_id="synthetic-duplicate",
        event_date="2024-02-03",
        installation_name="SYNTHETIC-ALIAS",
    )
    second["sources"][0].update(
        event_id="synthetic-duplicate", source_ref="synthetic-ref"
    )
    second["reviewed_text"] += [
        "synthetic-duplicate",
        "2024-02-03",
        "SYNTHETIC-ALIAS",
        "synthetic-ref",
    ]
    assert importer.ingest(second)["events_added"] == 0
    assert importer.ingest(second)["sources_added"] == 0
    manifest = importer.export(tmp_path / "export")
    assert manifest["event_count"] == 1 and manifest["source_count"] == 2
    log = json.loads((tmp_path / "export/dedup_log.json").read_text())
    assert {row["rule"] for row in log} == {"source_ref", "installation_date_country"}


def test_three_source_synthetic_worked_example(tmp_path, html, context, reviewed):
    """Demonstrates three inputs, not three verified sources for the real seed.

    OPRED seed date/quantity/cause: #1325 (2026). HSE issue/compliance dates and
    alert are synthetic. Expected answer is one event with three agency rows;
    no HSE mechanical cause or event date is inferred from its later notice.
    """
    packet = parse(html, context, reviewed)
    notice = load(workbook(tmp_path), reviewed).sources[0]
    alert = SourceDocument(
        event_id="synthetic-seed",
        source="iadc",
        source_ref="synthetic-alert",
        title="synthetic alert",
        doc_type="alert",
        issue_date=date(2024, 7, 24),
        status="published",
        url="https://iadc.org/safety-alerts/synthetic-example/",
        access_date=ACCESS,
        licence="All rights reserved; metadata and original paraphrase only",
        evidence_level="industry alert",
        access_status="public",
        association_basis="synthetic example association",
    )
    packet["sources"] += [notice.model_dump(mode="json"), alert.model_dump(mode="json")]
    importer = OffshoreImporter(tmp_path / "example.sqlite")
    result = importer.ingest(packet)
    assert result == {"events_added": 1, "sources_added": 3, "linked_existing": 0}
    manifest = importer.export(tmp_path / "example-export")
    assert manifest["event_count"] == 1 and manifest["source_count"] == 3
    assert manifest["bulk_admitted"] is False
    assert json.loads((tmp_path / "example-export/gates.json").read_text()) == {}


def test_truncated_workbook_returns_manual_step(tmp_path, reviewed):
    from zipfile import ZIP_DEFLATED, ZipFile

    path = workbook(tmp_path)
    broken = tmp_path / "truncated.xlsx"
    with ZipFile(path) as original, ZipFile(broken, "w", ZIP_DEFLATED) as output:
        for name in original.namelist():
            data = (
                b"<broken"
                if name == "xl/worksheets/sheet1.xml"
                else original.read(name)
            )
            output.writestr(name, data)
    result = load(broken, reviewed)
    assert result.status == "manual_required" and not result.sources


def test_mixed_oil_and_chemical_layout_is_not_summed(html, context, reviewed):
    extra = "<tr><th>Oil product released</th><td>synthetic oil</td></tr>"
    with pytest.raises(ValueError, match="PON1"):
        parse(
            html.replace("</main>", f"<table>{extra}</table></main>"),
            context,
            reviewed + ["synthetic oil"],
        )


def test_header_row_is_supplied_instead_of_guessed(tmp_path, reviewed):
    path = workbook(tmp_path)
    from openpyxl import load_workbook

    book = load_workbook(path)
    book["notices"].insert_rows(1)
    book["notices"].cell(1, 1, "synthetic workbook heading")
    book.save(path)
    book.close()
    assert load(path, reviewed, header_row=2).status == "ready"


def run_cli(arguments):
    import runpy

    script = Path(__file__).resolve().parents[3] / "scripts/import_hse_notices.py"
    return runpy.run_path(str(script))["main"](arguments)


def cli_config(tmp_path, context, reviewed):
    config = {
        "context": context,
        "reviewed_text": reviewed,
        "access_date": ACCESS.isoformat(),
        "pon1_url": PON1_URL,
        "columns": COLUMNS,
        "sheet": "notices",
        "header_row": 1,
        "associations": {"synthetic-notice": "synthetic-seed"},
        "date_format": "%d/%m/%Y",
        "sources": [
            {
                "event_id": "synthetic-seed",
                "source": "iadc",
                "source_ref": "synthetic-alert",
                "title": "synthetic alert",
                "doc_type": "alert",
                "issue_date": "2024-07-24",
                "status": "published",
                "url": "https://iadc.org/safety-alerts/synthetic-example/",
                "access_date": ACCESS.isoformat(),
                "licence": "All rights reserved; metadata and original paraphrase only",
                "evidence_level": "industry alert",
                "access_status": "public",
                "association_basis": "synthetic example association",
            }
        ],
    }
    path = tmp_path / "config.json"
    path.write_text(json.dumps(config))
    return path


def test_cli_three_source_example_and_rerun(tmp_path, context, reviewed, capsys):
    """Hand counts: one seed + OPRED/HSE/alert = one event, three sources.

    A repeated unique-key intake adds zero events and zero sources. The dates
    and quantity are fixture literals; the live third source is unverified.
    """
    config = cli_config(tmp_path, context, reviewed)
    root = tmp_path / "external-data"
    args = [
        "--pon1",
        str(FIXTURES / "pon1_synthetic.html"),
        "--hse-xlsx",
        str(workbook(tmp_path)),
        "--config",
        str(config),
        "--data-root",
        str(root),
    ]
    assert run_cli([*args, "--export", "first"]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["intake"]["sources_added"] == 3
    export = root / "modules/hse/offshore/first"
    manifest = json.loads((export / "manifest.json").read_text())
    assert manifest["event_count"] == 1 and manifest["source_count"] == 3
    assert manifest["bulk_admitted"] is False
    assert run_cli([*args, "--export", "second"]) == 0
    assert json.loads(capsys.readouterr().out)["intake"]["sources_added"] == 0


def test_cli_missing_spreadsheet_makes_no_database(tmp_path, context, reviewed, capsys):
    config = cli_config(tmp_path, context, reviewed)
    root = tmp_path / "external-data"
    args = [
        "--hse-xlsx",
        str(tmp_path / "absent.xlsx"),
        "--config",
        str(config),
        "--data-root",
        str(root),
        "--export",
        "first",
    ]
    assert run_cli(args) == 2
    result = json.loads(capsys.readouterr().out)
    assert HSE_SPREADSHEET_URL in result["manual_urls"]
    assert not root.exists()


def test_cli_manual_urls_need_no_credentials_or_configuration(capsys):
    assert run_cli(["--manual-hse"]) == 0
    result = json.loads(capsys.readouterr().out)
    assert result["reason"] == "BLOCKED-NETWORK"
    assert HSE_SPREADSHEET_URL in result["manual_urls"]


@pytest.mark.parametrize("name", ["..", "../escape", "/absolute"])
def test_cli_export_cannot_escape_external_root(tmp_path, context, reviewed, name):
    config = cli_config(tmp_path, context, reviewed)
    with pytest.raises(SystemExit):
        run_cli(
            [
                "--pon1",
                str(FIXTURES / "pon1_synthetic.html"),
                "--config",
                str(config),
                "--data-root",
                str(tmp_path / "external-data"),
                "--export",
                name,
            ]
        )


def test_cli_requires_external_data_root(tmp_path, context, reviewed):
    config = cli_config(tmp_path, context, reviewed)
    repo = Path(__file__).resolve().parents[3]
    with pytest.raises(SystemExit):
        run_cli(
            [
                "--pon1",
                str(FIXTURES / "pon1_synthetic.html"),
                "--config",
                str(config),
                "--data-root",
                str(repo / "data"),
                "--export",
                "first",
            ]
        )


def test_cli_refusal_does_not_echo_unreviewed_text(tmp_path, context, reviewed, capsys):
    config = cli_config(tmp_path, context, reviewed)
    content = json.loads(config.read_text())
    suspect = "Synthetic Individual"
    content["context"]["failure_summary"] = suspect
    config.write_text(json.dumps(content))
    assert (
        run_cli(
            [
                "--pon1",
                str(FIXTURES / "pon1_synthetic.html"),
                "--config",
                str(config),
                "--data-root",
                str(tmp_path / "external-data"),
                "--export",
                "first",
            ]
        )
        == 1
    )
    captured = capsys.readouterr()
    assert suspect not in captured.out + captured.err
    assert "refused" in captured.err
