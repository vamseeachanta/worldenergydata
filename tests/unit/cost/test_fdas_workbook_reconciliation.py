"""Read-only FDAS reconciliation, with independently calculated examples.

Workbook references: FDAS_V30 financial_project_summary.xlsx, Project_Summary;
lease_assumptions.xlsx, assumptions; generator generate_financial_summary_V30.py.
Fixture quantities are invented arithmetic inputs, never engineering defaults.
"""

import json
import os
import subprocess
import sys
import zipfile
from dataclasses import replace
from decimal import Decimal as D
from pathlib import Path

import pytest
from openpyxl import Workbook

from worldenergydata.cost.workbook_reconciliation import (
    Basis,
    CostValue,
    Interval,
    compare,
    freeze_inventory,
    read_cell,
    reconcile_workbooks,
    render_report,
    write_outputs,
)

ROOT = Path(__file__).resolve().parents[3]
FDAS = ROOT / "docs/modules/bsee/analysis/production/FDAS_V30"
CONTRACT = ROOT / "config/analysis/lower_tertiary/fdas_workbook_vintages.json"
BASIS = Basis(
    "USD", "nominal", "2018", "gross", "project", "capex", "USD", "2018", "FDAS_V30"
)


def cost(low="90", high="110", **changes):
    return CostValue(
        "example",
        "host",
        BASIS,
        Interval(D(low), D(high)),
        "modeled",
        "hand calculation",
        **changes,
    )


def test_interval_variance_and_explanation():
    """Hand calc: W=[100,120], C=[90,110] => [-10,30]. M=25 crosses
    the materiality boundary, so classification remains uncertain. Point
    W=150 minus C=100 gives 50 > 25, requiring supported explanation.
    """
    workbook = replace(cost("100", "120"), provenance="workbook_assumption")
    result = compare(workbook, cost(), D("25"))
    assert result["variance_low"] == "-10"
    assert result["variance_high"] == "30"
    assert result["materiality"] == "uncertain_materiality"
    assert result["status"] == "UNRESOLVED"
    result = compare(
        replace(workbook, amount=Interval(D("150"), D("150"))),
        cost("100", "100"),
        D("25"),
        "scope",
        "scope change record",
    )
    assert result["materiality"] == "material"
    assert result["status"] == "EXPLAINED"


@pytest.mark.parametrize(
    "field,value",
    [
        ("currency", "EUR"),
        ("price_basis", "real"),
        ("basis_year", "2020"),
        ("ownership", "net"),
        ("scope", "phase"),
        ("capex_basis", "opex"),
        ("unit", "USD MM"),
        ("period", "2019"),
        ("vintage", "FDAS_V50"),
    ],
)
def test_incompatible_basis_is_not_subtracted(field, value):
    result = compare(
        cost(), replace(cost(), basis=replace(BASIS, **{field: value})), D("1")
    )
    assert result["status"] == "NOT_EVALUATED"
    assert result["variance_low"] is None


def test_unknown_basis_and_missing_threshold_are_not_acceptance():
    unknown = replace(cost(), basis=replace(BASIS, basis_year=None))
    assert compare(unknown, unknown, D("1"))["status"] == "NOT_EVALUATED"
    assert compare(cost(), cost(), None)["materiality"] == "NOT_EVALUATED"


@pytest.mark.parametrize("low,high", [("NaN", "1"), ("0", "Infinity"), ("2", "1")])
def test_invalid_interval(low, high):
    with pytest.raises(ValueError):
        Interval(D(low), D(high))


def test_classification_and_evidence_required():
    with pytest.raises(ValueError, match="classification"):
        compare(cost(), cost(), D("1"), "other", "record")
    assert compare(cost(), cost(), D("1"), "scope", "")["status"] == "UNRESOLVED"
    assert (
        compare(cost(), cost(), D("1"), "missing evidence", "record")["status"]
        == "UNRESOLVED"
    )


def test_formula_cache_never_calculated_or_treated_as_zero(tmp_path):
    """openpyxl 3.1 load_workbook data_only returns stored cache, not evaluation."""
    path = tmp_path / "formula.xlsx"
    book = Workbook()
    book.active["A1"] = "=1+2"
    book.active["B1"] = 0
    book.save(path)
    cell = read_cell(path, "Sheet", "A1")
    assert cell == {"formula": "=1+2", "value": None, "status": "NOT_EVALUATED"}
    assert read_cell(path, "Sheet", "B1")["value"] == "0"


def test_formula_cached_value_is_unverified(tmp_path):
    """Hand calculation 1+2=3; a stored Excel cache is exposed but its
    freshness is not established (openpyxl 3.1 data_only documentation).
    """
    book = Workbook()
    book.active["A1"] = "=1+2"
    source = tmp_path / "source.xlsx"
    book.save(source)
    cached = tmp_path / "cached.xlsx"
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(cached, "w") as target:
        for name in archive.namelist():
            data = archive.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<v></v>", b"<v>3</v>")
            target.writestr(name, data)
    assert read_cell(cached, "Sheet", "A1") == {
        "formula": "=1+2",
        "value": "3",
        "status": "CACHED_UNVERIFIED",
    }


def test_frozen_inventory_and_schema_drift(tmp_path):
    contract = json.loads(CONTRACT.read_text())
    inventory = freeze_inventory(FDAS, contract)
    assert inventory["vintage"] == "FDAS_V30"
    assert len(inventory["workbooks"]["financial_project_summary.xlsx"]["sheets"]) == 11
    assert "creator" not in json.dumps(inventory)
    contract["workbooks"]["lease_assumptions.xlsx"]["sheets"]["assumptions"][
        "rows"
    ] += 1
    with pytest.raises(ValueError, match="schema drift"):
        freeze_inventory(FDAS, contract)
    contract = json.loads(CONTRACT.read_text())
    contract["workbooks"]["lease_assumptions.xlsx"]["sha256"] = "invalid"
    with pytest.raises(ValueError, match="digest mismatch"):
        freeze_inventory(FDAS, contract)


def test_component_crosswalk_and_v30_monthly_reproduction():
    """Project_Summary E3:I3 = 2000+0+0+0+630 = 2630 USD MM.
    assumptions!D15=100 USD MM; facilities=2730 USD MM (J3).
    K3+L3=965.6+821.7=1787.3 USD MM; total=4517.3 USD MM.
    Monthly host entries: 24 * 83333333.33333333 + 730000000
    = 2729999999.99999992; minus J3 yields -0.00000008 USD.
    No undocumented tolerance rounds this workbook residual to zero.
    These are workbook assumptions, not independent cost disclosures.
    """
    result = reconcile_workbooks(FDAS, json.loads(CONTRACT.read_text()))
    project = next(p for p in result["projects"] if p["project"] == "Big Foot")
    assert project["total_capex_usd"] == "4517300000"
    assert project["water_injection_facility_usd"] == "100000000"
    assert project["component_bridge_variance_usd"] == "0"
    assert D(project["monthly_checks"]["facilities"]["variance_usd"]) == D(
        "-0.00000008"
    )
    assert project["monthly_checks"]["facilities"]["status"] == "UNRESOLVED"
    # Workbook summaries use binary-float summation. Their tiny nonzero
    # residuals must remain visible until a sourced materiality policy exists.
    assert project["monthly_checks"]["drilling"]["status"] == "UNRESOLVED"
    assert project["monthly_checks"]["completion"]["status"] == "UNRESOLVED"
    assert {r["category"] for r in project["crosswalk"]} >= {
        "host",
        "SURF",
        "drilling",
        "completion",
        "pumps",
        "dry-system",
        "OPEX",
        "other",
    }
    assert all(
        r["disposition"] in {"included", "subtotal", "excluded"}
        for r in project["crosswalk"]
    )
    assert all(
        r["disclosed"] is None
        and r["modeled"] is None
        and r["mapped"] is None
        and r["variance"] is None
        for r in project["crosswalk"]
    )
    assert all(p["provenance"] == "workbook_assumption" for p in result["projects"])
    assert len(result["projects"]) == 10
    assert (
        next(p for p in result["projects"] if p["project"] == "Tiber")[
            "comparison_status"
        ]
        == "NOT_EVALUATED"
    )


def test_report_outputs_deterministic_and_escaped(tmp_path):
    result = reconcile_workbooks(FDAS, json.loads(CONTRACT.read_text()))
    first = tmp_path / "first"
    second = tmp_path / "second"
    before = {p.name: p.read_bytes() for p in FDAS.glob("*.xlsx")}
    write_outputs(result, first)
    write_outputs(result, second)
    assert {p.name: p.read_bytes() for p in first.iterdir()} == {
        p.name: p.read_bytes() for p in second.iterdir()
    }
    assert before == {p.name: p.read_bytes() for p in FDAS.glob("*.xlsx")}
    html = (first / "fdas_reconciliation.html").read_text()
    assert "<h2>Asset-to-total index</h2>" in html
    for text in (
        "Portfolio",
        "Revision history",
        "Asset-to-total",
        "Total-to-asset",
        "Time traces",
        "Estimator backtests",
        "Provenance",
        "maintainer-gated",
    ):
        assert text in html
    result["projects"][0]["project"] = "<script>alert(1)</script>"
    assert "<script>" not in render_report(result)
    manifest = json.loads((first / "fdas_reconciliation_manifest.json").read_text())
    assert set(manifest["outputs"]) == {
        "fdas_reconciliation.json",
        "fdas_reconciliation.csv",
        "fdas_reconciliation.html",
    }


def test_existing_cost_map_and_trace_are_evidence_not_workbook_equivalence():
    from worldenergydata.cost.workbook_reconciliation.integration import (
        attach_existing_evidence,
    )

    result = reconcile_workbooks(FDAS, json.loads(CONTRACT.read_text()))
    attach_existing_evidence(result, ROOT)
    assert result["upstream"]["cost_map"]["status"] == "PARTIAL"
    assert result["upstream"]["estimator"]["status"] == "NOT_EVALUATED"
    assert result["asset_evidence"]["project_id"] == "prj-000001"
    assert result["asset_evidence"]["requirements"]
    assert result["asset_evidence"]["targets"]
    assert result["asset_evidence"]["events"]
    project = next(p for p in result["projects"] if p["project"] == "Big Foot")
    assert project["project_id"] == "prj-000001"
    host = next(r for r in project["crosswalk"] if r["category"] == "host")
    assert host["requirement_ids"] == ["req-000001"]
    assert host["mapping_status"] == "mapped"
    assert all(p["comparison_status"] == "NOT_EVALUATED" for p in result["projects"])
    html = render_report(result)
    assert "prj-000001" in html
    assert "req-000001" in html


def test_output_collision_does_not_partially_write(tmp_path):
    result = reconcile_workbooks(FDAS, json.loads(CONTRACT.read_text()))
    existing = tmp_path / "fdas_reconciliation.json"
    existing.write_text("preserve")
    with pytest.raises(FileExistsError):
        write_outputs(result, tmp_path)
    assert existing.read_text() == "preserve"
    assert list(tmp_path.iterdir()) == [existing]


@pytest.mark.parametrize("threshold", [D("-1"), D("Infinity"), D("NaN")])
def test_invalid_threshold_rejected(threshold):
    with pytest.raises(ValueError, match="threshold"):
        compare(cost(), cost(), threshold)


def test_exact_threshold_boundaries_and_project_identity():
    """Hand calculation: [100,100]-[90,110]=[-10,10]; M=10
    includes both endpoints, while a different project cannot be subtracted.
    """
    assert compare(cost("100", "100"), cost(), D("10"))["status"] == "WITHIN_THRESHOLD"
    assert (
        compare(cost(), replace(cost(), project="another"), D("10"))["variance_low"]
        is None
    )
    assert (
        compare(cost(), replace(cost(), component="another"), D("10"))["variance_low"]
        is None
    )


def test_cli_produces_readable_pinned_outputs(tmp_path):
    completed = subprocess.run(
        [
            sys.executable,
            str(ROOT / "scripts/cost/build_fdas_workbook_reconciliation.py"),
            "--output",
            str(tmp_path / "output"),
        ],
        cwd=ROOT,
        env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"},
        text=True,
        capture_output=True,
        check=True,
    )
    assert "10 workbook projects" in completed.stdout
    manifest = json.loads(
        (tmp_path / "output/fdas_reconciliation_manifest.json").read_text()
    )
    assert manifest["producer_hashes"]
    assert manifest["evidence_input_hashes"]


def test_crosswalk_rejects_tampered_fingerprint(tmp_path):
    from worldenergydata.cost.workbook_reconciliation.integration import (
        _attach_crosswalk,
    )

    curated = ROOT / "data/modules/cost/curated"
    (tmp_path / "cost_project_identity.csv").write_bytes(
        (curated / "cost_project_identity.csv").read_bytes()
    )
    text = (curated / "fdas_project_cost_crosswalk.csv").read_text()
    original = json.loads(CONTRACT.read_text())["workbooks"][
        "financial_project_summary.xlsx"
    ]["sha256"]
    (tmp_path / "fdas_project_cost_crosswalk.csv").write_text(
        text.replace(original, "tampered")
    )
    result = reconcile_workbooks(FDAS, json.loads(CONTRACT.read_text()))
    with pytest.raises(ValueError, match="digest mismatch"):
        _attach_crosswalk(result, tmp_path)
