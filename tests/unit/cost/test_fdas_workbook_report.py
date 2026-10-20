"""Report arithmetic is checked by independent cell sums and synthetic examples."""

import csv
import hashlib
import json
from decimal import Decimal
from pathlib import Path

import pytest

from worldenergydata.cost.workbook_reconciliation.report import (
    build,
    compare_rows,
    render_html,
    workbook_rows,
)

ROOT = Path(__file__).resolve().parents[3]


@pytest.fixture(scope="module")
def rows():
    return workbook_rows(ROOT)[0]


def test_component_and_monthly_bridges(rows):
    """For each project, development = facilities + drilling + completion.

    Facilities = host + SURF + booster + injection pump + dry system +
    the explicitly identified injection-facility residual. This is an
    independent addition of the native workbook summary cells, not a model.
    """
    projects = {r["project_id"] for r in rows}
    assert len(projects) == 10  # Frozen financial workbook: summary rows 2:11.
    for project in projects:
        by_component = {r["component"]: r for r in rows if r["project_id"] == project}
        amounts = {
            key: Decimal(value["workbook_value"])
            for key, value in by_component.items()
            if value["workbook_value"]
        }
        num = amounts.__getitem__
        assert num("development_total") == num("facilities") + num("drilling") + num(
            "completion"
        )
        assert num("facilities") == sum(
            num(k)
            for k in [
                "host",
                "surf",
                "booster_pump",
                "injection_pump",
                "dry_system",
                "injection_facility",
            ]
        )
        assert num("dc_total") == num("drilling") + num("completion")
        for key in [
            "facilities",
            "drilling",
            "completion",
            "variable_opex",
            "fixed_opex",
        ]:
            r = by_component[key]
            assert Decimal(r["variance"]) == Decimal(r["workbook_value"]) - Decimal(
                r["monthly_value"]
            )
        assert by_component["installation"]["status"] == "NOT_EVALUATED"


def test_monthly_sums_independent_from_builder(rows):
    """Summary K2 is compared with the sum of monthly column I, independently."""
    from openpyxl import load_workbook

    w = load_workbook(
        ROOT
        / "docs/modules/bsee/analysis/production/FDAS_V30/financial_project_summary.xlsx",
        read_only=True,
        data_only=True,
    )
    try:
        title = w.worksheets[0]["A2"].value
        monthly = sum(
            Decimal(str(r[0]))
            for r in w[title].iter_rows(
                min_row=2, min_col=9, max_col=9, values_only=True
            )
        )
        observed = next(
            r
            for r in rows
            if r["project_id"] == "workbook-project-001"
            and r["component"] == "drilling"
        )
        assert Decimal(observed["monthly_value"]) == monthly
    finally:
        w.close()


def test_values_do_not_become_disclosures(rows):
    assert all(r["provenance"] == "workbook_assumption" for r in rows)
    assert all(r["disclosed_value"] == "" and r["modeled_value"] == "" for r in rows)
    assert all(r["materiality"] == "NOT_EVALUATED" for r in rows)


def test_escaping_and_sections():
    page = render_html(
        {
            "rows": [{"project_id": "<script>bad</script>", "component": "host"}],
            "inventory": {},
            "upstream": [],
            "assets": [],
            "traces": [],
            "comparisons": [],
        }
    )
    assert "<script>bad" not in page
    assert "&lt;script&gt;" in page
    for heading in [
        "Portfolio summary",
        "Project drill-down",
        "Asset to total",
        "Total to asset",
        "Cost time traces",
        "Estimator backtests",
        "Workbook reconciliation",
        "Provenance legend",
    ]:
        assert heading in page
    assert "#1017" in page and "#899" in page


def test_deterministic_build_and_hash_manifest(tmp_path):
    first, second = tmp_path / "first", tmp_path / "second"
    inputs = list(
        (ROOT / "docs/modules/bsee/analysis/production/FDAS_V30").glob("*.xlsx")
    )
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs}
    build(ROOT, first)
    build(ROOT, second)
    assert before == {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs}
    assert {p.name: p.read_bytes() for p in first.iterdir()} == {
        p.name: p.read_bytes() for p in second.iterdir()
    }
    manifest = json.loads((first / "manifest.json").read_text())
    for name, digest in manifest["outputs"].items():
        assert hashlib.sha256((first / name).read_bytes()).hexdigest() == digest
    assert all(entry["status"] == "NOT_EVALUATED" for entry in manifest["upstream"])
    with (first / "reconciliation.csv").open() as stream:
        assert (
            len(list(csv.DictReader(stream))) == 150
        )  # 15 explicit components per 10 frozen rows.


def test_missing_comparison_is_not_zero(rows):
    assert all(r["mapped_value"] == "" for r in rows)
    assert all(r["status"] == "NOT_EVALUATED" for r in rows)


@pytest.fixture
def external_spec():
    basis = {
        "project_id": "workbook-project-001",
        "component": "host",
        "start": "2024-01",
        "end": "2024-12",
        "currency": "USD",
        "unit": "USD",
        "assumption_vintage": "test",
        "price_basis": "nominal",
        "basis_year": "not_applicable",
        "ownership": "gross",
        "scope": "development",
        "capex_basis": "capex",
        "normalization_vintage": "native",
    }
    row = {**basis, "workbook_value": "100", "source": "synthetic:A1"}
    entry = {
        "project_id": basis["project_id"],
        "component": "host",
        "workbook_basis": basis,
        "basis_evidence": "synthetic:basis",
        "comparison": {
            "basis": dict(basis),
            "amount": {"low": "80", "high": "90"},
            "provenance": "mapped",
            "source": "synthetic:B1",
        },
    }
    return [row], {
        "comparisons": [entry],
        "policy": {
            "absolute": "25",
            "relative": "0",
            "source": "synthetic test policy",
            "approved": True,
        },
    }


def test_external_evidence_is_integrated(external_spec):
    """100-[80,90]=[10,20]; both endpoints lie inside +/-25."""
    rows, spec = external_spec
    result = compare_rows(rows, spec)[0]
    assert result["variance"] == {"low": "10", "high": "20"}
    assert result["status"] == "PASS"
    assert result["comparison"]["provenance"] == "mapped"


def test_external_identity_vintage_and_missing_evidence(external_spec):
    rows, spec = external_spec
    entry = spec["comparisons"][0]
    entry["comparison"]["basis"]["assumption_vintage"] = "alternative"
    assert compare_rows(rows, spec)[0]["variance"] is None
    entry["basis_evidence"] = ""
    assert compare_rows(rows, spec)[0]["status"] == "NOT_EVALUATED"
    entry["workbook_basis"]["unit"] = "USD MM"
    with pytest.raises(ValueError, match="frozen row"):
        compare_rows(rows, spec)


def test_duplicate_comparisons_rejected(external_spec):
    rows, spec = external_spec
    spec["comparisons"] *= 2
    with pytest.raises(ValueError, match="duplicate"):
        compare_rows(rows, spec)


def test_approval_is_not_a_truthy_string():
    from worldenergydata.cost.workbook_reconciliation.models import MaterialityPolicy

    with pytest.raises(TypeError, match="boolean"):
        MaterialityPolicy(Decimal(25), Decimal(0), "synthetic", "false")


def test_executable_worked_example(capsys):
    """Independent hand result: [100-90,120-80]=[10,40], crossing +/-25."""
    import runpy

    runpy.run_path(
        str(ROOT / "examples/fdas_workbook_reconciliation.py"), run_name="__main__"
    )
    assert (
        capsys.readouterr().out
        == "variance=[10,40]\nmateriality=uncertain_materiality; status=FAIL\n"
    )
