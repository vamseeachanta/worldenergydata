"""Synthetic financial regressions; no published workbook values are read."""

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from test_production_defects import REFERENCE, load_reference


@pytest.fixture
def finance():
    return load_reference("generate_financial_summary_V30")


@pytest.fixture
def inputs(finance, monkeypatch):
    """Synthetic rates/prices/costs are caller inputs, not field assumptions."""
    leases = pd.DataFrame(
        {"LEASE_NAME": ["A"], "DEV_NAME": ["development-a"], "DEV_SYSTEM": ["test"]}
    )
    assumptions = pd.DataFrame(
        {
            "DEV_SYSTEM": ["test"],
            "HOST_CAPEX_MM": [0],
            "INJECTORS_PER_PRODUCER": [0.5],
            "BOOSTER_PUMP_TRIGGER_PRODUCERS": [2],
            "BOOSTER_PUMP_MAX_COUNT": [3],
            "WATER_INJECTION_TRIGGER_PRODUCERS": [1],
            "WATER_INJECTION_PUMP_MAX_COUNT": [4],
            "MIRR_FINANCE_RATE": [(1.1**12) - 1],
            "MIRR_REINVEST_RATE": [(1.2**12) - 1],
        }
    )
    prod = pd.DataFrame(
        {
            "DEV_NAME": ["development-a"],
            "API_WELL_NUMBER": [101],
            "MONTH": pd.to_datetime(["2025-02-01"]),
            "MONTHLY_OIL_VOLUME": [10],
        }
    )
    dnc = pd.DataFrame(
        {
            "DEV_NAME": ["development-a"],
            "DEV_SYSTEM": ["test"],
            "LEASE_NAME": ["A"],
            "API_WELL_NUMBER": [101],
            "WELL_SPUD_DATE": pd.to_datetime(["2025-01-01"]),
            "TOTAL_DEPTH_DATE": pd.to_datetime(["2025-01-02"]),
            "DRILLING_DAYS": [2],
            "COMPLETION_DAYS": [0],
        }
    )
    prices = pd.DataFrame(
        {"MONTH": pd.to_datetime(["2025-01-01", "2025-02-01"]), "WTI_USD": [5, 5]}
    )
    for name, frame in (
        ("load_leases", leases),
        ("load_assumptions_fixed", assumptions),
        ("load_production", prod),
        ("load_dnc", dnc),
        ("load_wti", prices),
    ):
        monkeypatch.setattr(finance, name, lambda *args, frame=frame: frame.copy())
    return leases, assumptions, prod, dnc, prices


def generate(finance, tmp_path):
    out = tmp_path / "financial.xlsx"
    finance.generate("leases", "assumptions", "production", "days", "prices", out)
    return out


def test_b4_separate_finance_and_reinvestment_rates(finance):
    """Hand calculation, synthetic monthly flows [-100,50,-110,200].

    Microsoft MIRR function, web edition, Syntax/Remarks (retrieved 2026-10-22):
    FV+=50*1.2**2+200=272; abs(PV-)=100+110/1.1**2=2100/11.
    n=3, so MIRR=(748/525)**(1/3)-1; annual=(748/525)**4-1.
    Annual input rates are (1.1**12)-1 and (1.2**12)-1, giving the monthly
    rates above. Expected values use these rational expressions, independently
    of the implementation. See fdas_known_defects.html for textbook PV basis.
    """
    result = finance.excel_like_mirr(
        np.array([-100, 50, -110, 200]), (1.1**12) - 1, (1.2**12) - 1
    )
    assert result == pytest.approx((748 / 525) ** (1 / 3) - 1)


def test_b4_generator_uses_both_deck_rates(finance, inputs, monkeypatch, tmp_path):
    seen = []

    def mirr(cf, finance_rate, reinvest_rate):
        seen.append((finance_rate, reinvest_rate))
        return float("nan")

    monkeypatch.setattr(finance, "excel_like_mirr", mirr)
    generate(finance, tmp_path)
    assert seen == [((1.1**12) - 1, (1.2**12) - 1)]


def test_b4_missing_rate_requires_input(finance, inputs, tmp_path):
    inputs[1].drop(columns="MIRR_FINANCE_RATE", inplace=True)
    with pytest.raises(ValueError, match="MIRR_FINANCE_RATE"):
        generate(finance, tmp_path)


def test_b7_outside_host_window_raises(finance, inputs, tmp_path):
    inputs[1]["HOST_CAPEX_MM"] = 1
    inputs[1]["HOST_PREFO_MONTHS"] = 2
    with pytest.raises(ValueError, match="CAPEX"):
        generate(finance, tmp_path)


def test_b7_full_host_window_conserves_capex(finance, inputs, tmp_path):
    """Synthetic one-month spread: supplied 1 million USD remains 1 million."""
    inputs[1]["HOST_CAPEX_MM"] = 1
    inputs[1]["HOST_PREFO_MONTHS"] = 1
    out = generate(finance, tmp_path)
    result = pd.read_excel(out, sheet_name="Project_Summary")
    assert result["Host CAPEX USD"].tolist() == [1_000_000]


def test_b10_missing_month_raises(finance, inputs, tmp_path):
    inputs[4].drop(index=1, inplace=True)
    with pytest.raises(ValueError, match="WTI"):
        generate(finance, tmp_path)


def test_b12_duplicate_price_month_raises(finance, inputs, tmp_path):
    inputs[4].loc[2] = inputs[4].iloc[0]
    with pytest.raises(pd.errors.MergeError):
        generate(finance, tmp_path)


def test_b5_emitted_header_matches_golden_header_only(finance, inputs, tmp_path):
    """Schema is the shipped header, not an assertion about golden values."""
    out = generate(finance, tmp_path)

    def header(path):
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            return list(
                next(wb["Project_Summary"].iter_rows(max_row=1, values_only=True))
            )
        finally:
            wb.close()

    assert header(out) == header(REFERENCE / "financial_project_summary.xlsx")


def test_b14_named_totals_excluded_and_unnamed_well_retained(finance, monkeypatch):
    """Valid API 101 contributes 2 days; a named total lacking API contributes none."""
    data = pd.DataFrame(
        {
            "LEASE_NAME": [None, "TOTAL"],
            "API_WELL_NUMBER": [101, None],
            "WELL_SPUD_DATE": ["2025-01-01", None],
            "DRILLING_DAYS": [2, 2],
            "COMPLETION_DAYS": [0, 0],
        }
    )
    monkeypatch.setattr(finance, "read_xlsx_any", lambda fp: [("days", data)])
    leases = pd.DataFrame(
        {"LEASE_NAME": ["A"], "DEV_NAME": ["development-a"], "DEV_SYSTEM": ["test"]}
    )
    result = finance.load_dnc("unused", leases)
    assert result["API_WELL_NUMBER"].tolist() == [101]
    assert result["DRILLING_DAYS"].sum() == 2


def test_b11_universe_diagnostic_does_not_add_developments(finance):
    """Two drilled developments minus one producing development leaves one."""
    prod = pd.DataFrame({"DEV_NAME": ["A"], "MONTHLY_OIL_VOLUME": [10]})
    days = pd.DataFrame({"DEV_NAME": ["A", "B"], "API_WELL_NUMBER": [101, 102]})
    result = finance.diagnose_universe(prod, days)
    assert result == {"status": "NOT_EVALUATED", "drilled_without_production": 1}


@pytest.mark.parametrize("rate", [-1, float("nan"), float("inf")])
def test_b4_invalid_rates_rejected(finance, rate):
    with pytest.raises(ValueError, match="MIRR"):
        finance.excel_like_mirr(np.array([-100, 110]), rate, 0)


def test_b4_nonfinite_cashflow_rejected(finance):
    with pytest.raises(ValueError, match="cashflow"):
        finance.excel_like_mirr(np.array([-100, float("nan"), 110]), 0, 0)


def test_b11_diagnostic_export_preserves_existing_universe(finance, inputs, tmp_path):
    other = inputs[3].iloc[0].copy()
    other["DEV_NAME"], other["API_WELL_NUMBER"] = "unproduced", 102
    inputs[3].loc[1] = other
    out = generate(finance, tmp_path)
    assert pd.read_excel(out, sheet_name="Project_Summary")[
        "Project Name"
    ].tolist() == ["development-a"]
    assert pd.read_excel(out, sheet_name="Diagnostics")[
        "drilled_without_production"
    ].tolist() == [1]


def test_b12_financial_lease_conflicts_raise(finance, monkeypatch):
    frame = pd.DataFrame(
        {
            "LEASE_NAME": ["A", "A"],
            "DEV_NAME": ["one", "two"],
            "DEV_SYSTEM": ["test", "test"],
        }
    )
    monkeypatch.setattr(finance, "read_xlsx_any", lambda fp: [("leases", frame)])
    with pytest.raises(pd.errors.MergeError):
        finance.load_leases("unused")


COUNT_ASSUMPTIONS = [
    "INJECTORS_PER_PRODUCER",
    "BOOSTER_PUMP_TRIGGER_PRODUCERS",
    "BOOSTER_PUMP_MAX_COUNT",
    "WATER_INJECTION_TRIGGER_PRODUCERS",
    "WATER_INJECTION_PUMP_MAX_COUNT",
]


def six_producers(inputs):
    """Six synthetic producing APIs keep the physical input independent of counts."""
    prod = inputs[2]
    for i in range(1, 6):
        prod.loc[i] = prod.iloc[0]
        prod.loc[i, "API_WELL_NUMBER"] = 101 + i


@pytest.mark.parametrize("name", COUNT_ASSUMPTIONS)
def test_reference_count_missing_assumption_rejected(finance, inputs, tmp_path, name):
    six_producers(inputs)
    inputs[1].drop(columns=name, inplace=True)
    with pytest.raises(ValueError, match=name):
        generate(finance, tmp_path)
    assert not (tmp_path / "financial.xlsx").exists()


@pytest.mark.parametrize("name", COUNT_ASSUMPTIONS)
@pytest.mark.parametrize("value", [None, "invalid", float("nan"), float("inf"), -1])
def test_reference_count_invalid_assumption_rejected(
    finance, inputs, tmp_path, name, value
):
    six_producers(inputs)
    inputs[1][name] = value
    with pytest.raises(ValueError, match=name):
        generate(finance, tmp_path)
    assert not (tmp_path / "financial.xlsx").exists()


@pytest.mark.parametrize("name", COUNT_ASSUMPTIONS[1:])
def test_reference_count_fractional_integer_rejected(finance, inputs, tmp_path, name):
    six_producers(inputs)
    inputs[1][name] = 1.5
    with pytest.raises(ValueError, match=name):
        generate(finance, tmp_path)


@pytest.mark.parametrize(
    "name", ["BOOSTER_PUMP_TRIGGER_PRODUCERS", "WATER_INJECTION_TRIGGER_PRODUCERS"]
)
def test_reference_count_zero_trigger_rejected(finance, inputs, tmp_path, name):
    six_producers(inputs)
    inputs[1][name] = 0
    with pytest.raises(ValueError, match=name):
        generate(finance, tmp_path)


@pytest.mark.parametrize("booster_cap, water_cap", [(3, 4), (1, 1), (0, 0)])
def test_reference_counts_and_costs_follow_supplied_deck(
    finance, inputs, tmp_path, booster_cap, water_cap
):
    """Hand calculation with synthetic supplied rules and unit costs.

    Injectors=floor(6*0.5)=3; uncapped boosters=floor(6/2)=3;
    uncapped water pumps=floor(6/1)=6. Caps (3,4), (1,1), (0,0) therefore
    give counts (3,4), (1,1), (0,0). At 2 and 3 million USD per pump,
    corresponding costs are (6,12), (2,3), (0,0) million USD.
    These inputs illustrate caller-supplied rules, not qualified design values.
    """
    six_producers(inputs)
    assumptions = inputs[1]
    assumptions["BOOSTER_PUMP_MAX_COUNT"] = booster_cap
    assumptions["WATER_INJECTION_PUMP_MAX_COUNT"] = water_cap
    assumptions["BOOSTER_PUMP_15K_MM"] = 2
    assumptions["WATER_INJECTION_PUMP_MM"] = 3
    expected = {
        (3, 4): (3, 4, 6_000_000, 12_000_000),
        (1, 1): (1, 1, 2_000_000, 3_000_000),
        (0, 0): (0, 0, 0, 0),
    }[(booster_cap, water_cap)]
    result = pd.read_excel(generate(finance, tmp_path), sheet_name="Project_Summary")
    assert result["PRODUCER WELLS"].tolist() == [6]
    assert result["INJECTOR WELLS"].tolist() == [3]
    for column, value in zip(
        [
            "Booster Pump Count",
            "Water Injection Pump Count",
            "Booster Pump USD",
            "Water Injection Pump USD",
        ],
        expected,
    ):
        assert result[column].tolist() == pytest.approx([value])


def test_reference_count_dry_system_needs_no_pump_assumptions(
    finance, inputs, tmp_path
):
    for frame in (inputs[0], inputs[1], inputs[3]):
        frame["DEV_SYSTEM"] = "dry"
    inputs[1].drop(columns=COUNT_ASSUMPTIONS[1:], inplace=True)
    result = pd.read_excel(generate(finance, tmp_path), sheet_name="Project_Summary")
    assert result["Booster Pump Count"].tolist() == [0]
    assert result["Water Injection Pump Count"].tolist() == [0]


def test_reference_count_no_injectors_needs_no_water_pump_assumptions(
    finance, inputs, tmp_path
):
    inputs[1]["INJECTORS_PER_PRODUCER"] = 0
    inputs[1].drop(columns=COUNT_ASSUMPTIONS[3:], inplace=True)
    result = pd.read_excel(generate(finance, tmp_path), sheet_name="Project_Summary")
    assert result["INJECTOR WELLS"].tolist() == [0]
    assert result["Water Injection Pump Count"].tolist() == [0]
