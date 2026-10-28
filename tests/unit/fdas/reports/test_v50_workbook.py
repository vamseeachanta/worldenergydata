"""V50 export contract; synthetic amounts are independent hand calculations."""

from contextlib import closing
from dataclasses import replace
from datetime import date, datetime, timezone

import pandas as pd
import pytest
from openpyxl import load_workbook

from worldenergydata.fdas.reports import v50_workbook


@pytest.fixture
def metadata():
    return v50_workbook.V50WorkbookMetadata(
        assumption_version="test-basis-unaccepted",
        git_commit="a" * 40,
        production_cutoff=date(2026, 4, 30),
        wti_deck_source="Synthetic test input; no market prices",
        npv_reference_convention="First material expenditure; annual period zero",
        generated_at=datetime(2026, 10, 19, tzinfo=timezone.utc),
    )


@pytest.fixture
def summary():
    """Hand calculation, USD: -100 at t=0, +121 at t=1, r=10%.

    Pre-tax NPV = -100 + 121/1.1 = 10. After-tax receipt assumed 110:
    after-tax NPV = -100 + 110/1.1 = 0. These are synthetic annual
    examples, not model assumptions; the exporter must not recalculate them.
    All other numeric cells use the synthetic input zero, not missing values.
    Schema source: repository FDAS_V30/financial_project_summary.xlsx,
    Project_Summary header row; additions follow issue 901 metadata scope.
    """
    row = dict.fromkeys(v50_workbook.LEGACY_SUMMARY_COLUMNS, 0)
    row.update(
        {
            "Project Name": "Synthetic project",
            "DEV SYSTEM": "Synthetic system",
            "FO Month": date(2026, 1, 1),
            "NPV USD": 10.0,
            "After-tax NPV USD": 0.0,
            "NPV Reference Date": date(2025, 1, 1),
        }
    )
    return pd.DataFrame([row])


def test_summary_and_metadata_roundtrip(tmp_path, summary, metadata):
    target = tmp_path / "summary.xlsx"
    original = summary.copy(deep=True)
    assert v50_workbook.write_v50_workbook(target, summary, metadata) == target
    with closing(load_workbook(target, data_only=True)) as book:
        assert book.sheetnames == ["Project_Summary", "Metadata"]
        rows = list(book["Project_Summary"].values)
        assert rows[0] == v50_workbook.SUMMARY_COLUMNS
        result = dict(zip(rows[0], rows[1]))
        assert result["NPV USD"] == 10.0
        assert result["After-tax NPV USD"] == 0.0  # Real zero, not missing.
        assert result["NPV Reference Date"] == datetime(2025, 1, 1)
        info = dict(book["Metadata"].values)
        assert info["assumption_version"] == metadata.assumption_version
        assert info["git_commit"] == metadata.git_commit
        assert info["npv_reference_convention"] == metadata.npv_reference_convention
        assert info["production_cutoff"] == "2026-04-30"
        assert info["wti_deck_source"] == metadata.wti_deck_source
        assert info["generated_at"] == "2026-10-19T00:00:00+00:00"
        assert info["canonical_acceptance_status"] == "NOT_EVALUATED"
        assert info["npv_units"] == "USD"
        assert info["pre_tax_npv_column"] == "NPV USD"
        assert info["after_tax_npv_column"] == "After-tax NPV USD"
        assert info["lineage"] == "V30 is lineage-only"
    pd.testing.assert_frame_equal(summary, original)


def test_schema_matches_frozen_workbook():
    """The committed lineage workbook is the independent schema reference."""
    from pathlib import Path

    root = Path(__file__).resolve().parents[4]
    source = (
        root
        / "docs/modules/bsee/analysis/production/FDAS_V30/financial_project_summary.xlsx"
    )
    with closing(load_workbook(source, read_only=True)) as book:
        assert (
            tuple(next(book["Project_Summary"].values))
            == v50_workbook.LEGACY_SUMMARY_COLUMNS
        )


@pytest.mark.parametrize(
    "column", ["NPV USD", "After-tax NPV USD", "NPV Reference Date"]
)
def test_missing_required_column_rejected(tmp_path, summary, metadata, column):
    with pytest.raises(ValueError, match="Missing"):
        v50_workbook.write_v50_workbook(
            tmp_path / "bad.xlsx", summary.drop(columns=column), metadata
        )
    assert not (tmp_path / "bad.xlsx").exists()


@pytest.mark.parametrize(
    "value", [float("nan"), float("inf"), None, "NOT_EVALUATED", True]
)
def test_missing_or_invalid_npv_rejected(tmp_path, summary, metadata, value):
    summary["After-tax NPV USD"] = value
    with pytest.raises(ValueError, match="finite numeric"):
        v50_workbook.write_v50_workbook(tmp_path / "bad.xlsx", summary, metadata)
    assert not (tmp_path / "bad.xlsx").exists()


@pytest.mark.parametrize(
    "field,value",
    [
        ("assumption_version", " "),
        ("wti_deck_source", ""),
        ("npv_reference_convention", ""),
        ("git_commit", "unversioned"),
        ("production_cutoff", "2026-04-30"),
        ("generated_at", datetime(2026, 10, 19)),
    ],
)
def test_invalid_provenance_rejected(tmp_path, summary, metadata, field, value):
    with pytest.raises(ValueError, match=field):
        v50_workbook.write_v50_workbook(
            tmp_path / "bad.xlsx", summary, replace(metadata, **{field: value})
        )
    assert not (tmp_path / "bad.xlsx").exists()


def test_not_applicable_lineage_pattern_roundtrip(tmp_path, summary, metadata):
    """Synthetic rows mirror the blank-cell pattern of lineage rows 7 through 11."""
    summary = pd.concat([summary] * 5, ignore_index=True)
    summary["Project Name"] = [f"Synthetic row {row}" for row in range(7, 12)]
    oil_reason = "No first oil by production cutoff"
    mirr_reason = "MIRR undefined: cash flows lack both signs"
    for row in (0, 1, 4):
        summary.loc[row, "FO Month"] = v50_workbook.NotApplicable(oil_reason)
    for column in ("MIRR monthly", "MIRR annual"):
        summary[column] = summary[column].astype(object)
        for row in (0, 1, 2, 4):
            summary.loc[row, column] = v50_workbook.NotApplicable(mirr_reason)
    # Metadata addresses refer to exported positions, independent of input order/index.
    summary = summary.loc[:, list(reversed(summary.columns))]
    summary.index = [20, 20, 30, 40, 50]
    original = summary.copy(deep=True)
    target = tmp_path / "summary.xlsx"
    v50_workbook.write_v50_workbook(target, summary, metadata)
    with closing(load_workbook(target, data_only=True)) as book:
        sheet = book["Project_Summary"]
        assert sheet.max_row == 6
        assert [sheet.cell(row, 1).value for row in range(2, 7)] == [
            f"Synthetic row {row}" for row in range(7, 12)
        ]
        expected = {
            "C2": oil_reason,
            "C3": oil_reason,
            "C6": oil_reason,
            "T2": mirr_reason,
            "T3": mirr_reason,
            "T4": mirr_reason,
            "T6": mirr_reason,
            "U2": mirr_reason,
            "U3": mirr_reason,
            "U4": mirr_reason,
            "U6": mirr_reason,
        }
        info = dict(book["Metadata"].values)
        reasons = {
            key: value
            for key, value in info.items()
            if key.startswith("not_applicable:")
        }
        assert reasons == {
            f"not_applicable:Project_Summary!{cell}": reason
            for cell, reason in expected.items()
        }
        for cell in expected:
            assert sheet[cell].value is None
        assert sheet["C4"].value == datetime(2026, 1, 1)
        assert sheet["T5"].value == 0
        assert sheet["U5"].value == 0
    pd.testing.assert_frame_equal(summary, original)


@pytest.mark.parametrize("reason", [None, "", " ", 0])
def test_not_applicable_requires_reason(tmp_path, summary, metadata, reason):
    with pytest.raises(ValueError, match="reason.*nonblank text"):
        summary["MIRR monthly"] = v50_workbook.NotApplicable(reason)
        v50_workbook.write_v50_workbook(tmp_path / "bad.xlsx", summary, metadata)
    assert not (tmp_path / "bad.xlsx").exists()


@pytest.mark.parametrize(
    "column", ["Project Name", "NPV USD", "After-tax NPV USD", "NPV Reference Date"]
)
def test_not_applicable_rejected_outside_nullable_columns(
    tmp_path, summary, metadata, column
):
    summary[column] = v50_workbook.NotApplicable("Not applicable")
    with pytest.raises(ValueError, match="not nullable"):
        v50_workbook.write_v50_workbook(tmp_path / "bad.xlsx", summary, metadata)
    assert not (tmp_path / "bad.xlsx").exists()


@pytest.mark.parametrize("column", ["FO Month", "MIRR monthly", "MIRR annual"])
@pytest.mark.parametrize("value", [None, float("nan"), pd.NaT])
def test_nullable_columns_still_reject_unexplained_missing_values(
    tmp_path, summary, metadata, column, value
):
    summary[column] = pd.Series([value], dtype=object)
    with pytest.raises(ValueError, match="must be a date|finite numeric"):
        v50_workbook.write_v50_workbook(tmp_path / "bad.xlsx", summary, metadata)
    assert not (tmp_path / "bad.xlsx").exists()


@pytest.mark.parametrize(
    "case,message",
    [
        ("duplicate column", "columns must be unique"),
        ("unknown column", "Unexpected summary columns"),
        ("blank text", "DEV SYSTEM must be nonblank text"),
        ("invalid date", "NPV Reference Date must be a date"),
        ("NaT date", "NPV Reference Date must be a date"),
        ("timezone date", "NPV Reference Date must not contain a timezone"),
    ],
)
def test_summary_rejection_guards(tmp_path, summary, metadata, case, message):
    if case == "duplicate column":
        summary = pd.concat([summary, summary[["NPV USD"]]], axis=1)
    elif case == "unknown column":
        summary["Unmapped result"] = 1
    elif case == "blank text":
        summary["DEV SYSTEM"] = " "
    else:
        summary["NPV Reference Date"] = {
            "invalid date": "2025-01-01",
            "NaT date": pd.NaT,
            "timezone date": datetime(2025, 1, 1, tzinfo=timezone.utc),
        }[case]
    with pytest.raises(ValueError, match=message):
        v50_workbook.write_v50_workbook(tmp_path / "bad.xlsx", summary, metadata)
    assert not (tmp_path / "bad.xlsx").exists()


def test_remaining_economics_have_explicit_tax_basis(tmp_path, summary, metadata):
    target = tmp_path / "summary.xlsx"
    v50_workbook.write_v50_workbook(target, summary, metadata)
    with closing(load_workbook(target, data_only=True)) as book:
        info = dict(book["Metadata"].values)
        for column in ("Net CashFlow USD", "MIRR monthly", "MIRR annual"):
            assert info[f"tax_basis:{column}"] == "pre-tax"


def test_not_applicable_reason_is_literal_text(tmp_path, summary, metadata):
    summary["MIRR monthly"] = v50_workbook.NotApplicable("=1+1")
    target = tmp_path / "summary.xlsx"
    v50_workbook.write_v50_workbook(target, summary, metadata)
    with closing(load_workbook(target)) as book:
        info = {row[0].value: row[1] for row in book["Metadata"]}
        reason = info["not_applicable:Project_Summary!T2"]
        assert reason.value == "=1+1"
        assert reason.data_type == "s"


def test_duplicate_projects_rejected(tmp_path, summary, metadata):
    with pytest.raises(ValueError, match="unique"):
        v50_workbook.write_v50_workbook(
            tmp_path / "bad.xlsx", pd.concat([summary, summary]), metadata
        )


def test_empty_summary_rejected(tmp_path, summary, metadata):
    with pytest.raises(ValueError, match="empty"):
        v50_workbook.write_v50_workbook(
            tmp_path / "bad.xlsx", summary.iloc[:0], metadata
        )


def test_formula_like_text_is_stored_as_text(tmp_path, summary, metadata):
    summary.loc[0, "Project Name"] = "=1+1"
    target = tmp_path / "summary.xlsx"
    v50_workbook.write_v50_workbook(
        target, summary, replace(metadata, wti_deck_source="=1+1")
    )
    with closing(load_workbook(target)) as book:
        assert book["Project_Summary"]["A2"].data_type == "s"
        info = {row[0].value: row[1] for row in book["Metadata"]}
        assert info["wti_deck_source"].data_type == "s"


def test_existing_output_is_preserved(tmp_path, summary, metadata):
    target = tmp_path / "summary.xlsx"
    target.write_bytes(b"existing evidence")
    with pytest.raises(FileExistsError):
        v50_workbook.write_v50_workbook(target, summary, metadata)
    assert target.read_bytes() == b"existing evidence"
